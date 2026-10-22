"""Teacher XLSX parsing, validation and preview construction."""

import re
import secrets
import time
from openpyxl import load_workbook
from essay_platform.core.config import CLASSES_PATH, USERS_PATH
from essay_platform.persistence.datasets import load_json
from essay_platform.services.directory import clean_cell, validate_employee_no
from essay_platform.services.import_rules import normalize_header, parse_class_number, parse_combined_class, parse_grade_number

TEACHER_IMPORT_HEADERS = {
    "display_name": {"姓名", "name", "教师姓名"},
    "employeeNo": {"工号", "employeeno", "employee_id", "教师工号"},
    "subject": {"学科", "subject", "任教学科"},
    "classes": {"班级", "class", "classes", "任教班级"},
}
def find_teacher_import_header(rows):
    normalized_aliases = {
        key: {normalize_header(alias) for alias in aliases}
        for key, aliases in TEACHER_IMPORT_HEADERS.items()
    }
    for index, row in enumerate(rows[:20]):
        normalized = [normalize_header(value) for value in row]
        mapping = {}
        for field, aliases in normalized_aliases.items():
            match = next((position for position, value in enumerate(normalized)
                          if value in aliases), None)
            if match is not None:
                mapping[field] = match
        if len(mapping) == len(TEACHER_IMPORT_HEADERS):
            return index, mapping
    raise ValueError("未找到教师表头，请使用：姓名、工号、学科、班级")

def resolve_teacher_class_token(token, classes):
    compact = re.sub(r"\s+", "", clean_cell(token)).lower()
    exact = [item for item in classes if compact in {
        re.sub(r"\s+", "", str(item.get("id", ""))).lower(),
        re.sub(r"\s+", "", str(item.get("name", ""))).lower(),
    }]
    if len(exact) == 1:
        return exact[0]
    combined = parse_combined_class(token)
    if combined:
        grade_number, class_number = combined
        matches = [item for item in classes
                   if parse_grade_number(item.get("grade")) == grade_number
                   and parse_class_number(item.get("klass")) == class_number]
        if len(matches) == 1:
            return matches[0]
    return None

def parse_teacher_import_file(file):
    if load_workbook is None:
        raise ValueError("服务器缺少 openpyxl，无法读取 Excel")
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise ValueError("教师导入只支持 .xlsx 文件")
    try:
        file.stream.seek(0)
        workbook = load_workbook(file.stream, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Excel 文件无法读取：{exc}") from exc
    try:
        selected = None
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            try:
                header_index, mapping = find_teacher_import_header(rows)
            except ValueError:
                continue
            selected = worksheet.title, rows, header_index, mapping
            break
        if not selected:
            raise ValueError("所有工作表中都未找到教师表头")
        sheet_name, rows, header_index, mapping = selected
        parsed_rows = []
        for row_number, values in enumerate(rows[header_index + 1:], header_index + 2):
            if not any(clean_cell(value) for value in values):
                continue
            def cell(field):
                position = mapping[field]
                return clean_cell(values[position] if position < len(values) else "")
            parsed_rows.append({
                "rowNo": row_number,
                "display_name": cell("display_name"),
                "employeeNo": cell("employeeNo"),
                "subject": cell("subject"),
                "classesRaw": cell("classes"),
            })
            if len(parsed_rows) > 2000:
                raise ValueError("教师 Excel 最多允许 2000 行")
        if not parsed_rows:
            raise ValueError("Excel 中没有教师数据")
        return {"sheetName": sheet_name, "headerRow": header_index + 1,
                "rows": parsed_rows}
    finally:
        workbook.close()

def build_teacher_import_preview(file):
    parsed = parse_teacher_import_file(file)
    users = load_json(USERS_PATH, {"users": []}).get("users", [])
    classes = load_json(CLASSES_PATH, {"classes": []}).get("classes", [])
    teachers_by_employee = {
        (user.get("employeeNo") or user.get("username")): user
        for user in users if user.get("role") == "teacher"
    }
    users_by_username = {user.get("username"): user for user in users}
    employee_counts = {}
    for source in parsed["rows"]:
        employee_no = source.get("employeeNo", "")
        employee_counts[employee_no] = employee_counts.get(employee_no, 0) + 1
    preview_rows = []
    for source in parsed["rows"]:
        errors = []
        employee_no = source["employeeNo"]
        if not source["display_name"]:
            errors.append("姓名不能为空")
        if not validate_employee_no(employee_no):
            errors.append("工号格式无效")
        if not source["subject"]:
            errors.append("学科不能为空")
        if employee_no and employee_counts.get(employee_no, 0) > 1:
            errors.append("Excel 中工号重复")
        class_tokens = [item.strip() for item in re.split(r"[、，,；;\r\n]+", source["classesRaw"])
                        if item.strip()]
        if not class_tokens:
            errors.append("班级不能为空")
        class_ids, class_names = [], []
        for token in class_tokens:
            matched = resolve_teacher_class_token(token, classes)
            if not matched:
                errors.append(f"班级不存在：{token}")
                continue
            if matched["id"] not in class_ids:
                class_ids.append(matched["id"])
                class_names.append(matched.get("name") or matched["id"])
        existing = teachers_by_employee.get(employee_no)
        occupied = users_by_username.get(employee_no)
        if occupied and occupied.get("role") != "teacher":
            errors.append("工号已被非老师登录账号占用")
        action = "update" if existing else "new"
        preview_rows.append({
            **source,
            "username": existing.get("username") if existing else employee_no,
            "classIds": class_ids,
            "classNames": class_names,
            "action": action,
            "status": "error" if errors else action,
            "errors": errors,
        })
    summary = {
        "total": len(preview_rows),
        "new": sum(row["status"] == "new" for row in preview_rows),
        "update": sum(row["status"] == "update" for row in preview_rows),
        "error": sum(row["status"] == "error" for row in preview_rows),
        "sheetName": parsed["sheetName"],
        "headerRow": parsed["headerRow"],
    }
    return {"token": secrets.token_hex(16), "kind": "teacher", "createdAt": time.time(),
            "summary": summary, "rows": preview_rows}

def validate_teacher_import_preview_fresh(preview):
    users = load_json(USERS_PATH, {"users": []}).get("users", [])
    classes = {item["id"] for item in load_json(CLASSES_PATH, {"classes": []}).get("classes", [])}
    teachers_by_employee = {
        (user.get("employeeNo") or user.get("username")): user
        for user in users if user.get("role") == "teacher"
    }
    users_by_username = {user.get("username"): user for user in users}
    errors = []
    for row in preview.get("rows", []):
        employee_no = row.get("employeeNo", "")
        existing = teachers_by_employee.get(employee_no)
        current_action = "update" if existing else "new"
        if current_action != row.get("action"):
            errors.append(f"第 {row.get('rowNo')} 行：老师账号数据已变化")
        elif existing and existing.get("username") != row.get("username"):
            errors.append(f"第 {row.get('rowNo')} 行：工号已指向其他账号")
        elif not existing and users_by_username.get(employee_no):
            errors.append(f"第 {row.get('rowNo')} 行：工号登录账号已被占用")
        missing = [class_id for class_id in row.get("classIds", []) if class_id not in classes]
        if missing:
            errors.append(f"第 {row.get('rowNo')} 行：班级数据已变化")
    return errors
