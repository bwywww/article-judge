"""Student XLSX parsing, validation and preview construction."""

import re
import secrets
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from itertools import islice
from openpyxl import load_workbook
from essay_platform.core.auth import secure_password_hash, validate_new_password
from essay_platform.core.config import (
    CLASSES_PATH, DEFAULT_TEMP_PASSWORD, MAX_EXCEL_ROWS,
    MAX_EXCEL_UNCOMPRESSED_BYTES, USERS_PATH,
)
from essay_platform.persistence.datasets import load_json
from essay_platform.services.directory import (
    class_name, clean_cell, validate_student_no, validate_username,
)
from essay_platform.services.import_rules import (
    HEADER_LOOKUP, generated_student_no, grade_year_in_stage,
    import_class_sort_key, normalize_header, normalize_import_class,
    parse_academic_year_start, parse_class_number,
)


def find_import_header(rows):
    best = None
    for idx, row in enumerate(rows[:8]):
        mapping = {}
        for col, cell in enumerate(row):
            key = HEADER_LOOKUP.get(normalize_header(cell))
            if key and key not in mapping:
                mapping[key] = col
        if "display_name" not in mapping or not ({"grade", "klass", "studentNo"} & set(mapping)):
            continue
        score = len(set(mapping) & {"grade", "klass", "studentNo", "display_name"})
        if not best or score > best[0]:
            best = (score, idx, mapping)
    if not best or best[0] < 2:
        return None, {}
    return best[1], best[2]

def parse_student_import_file(file, academic_year_start=None):
    if load_workbook is None:
        raise ValueError("缺少 Excel 依赖，请先安装 openpyxl")
    academic_year_start = parse_academic_year_start(academic_year_start)
    filename = file.filename or ""
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("请上传 .xlsx 格式的 Excel 文件")
    try:
        file.stream.seek(0)
        with zipfile.ZipFile(file.stream) as archive:
            expanded_size = sum(item.file_size for item in archive.infolist())
            if expanded_size > MAX_EXCEL_UNCOMPRESSED_BYTES:
                raise ValueError(
                    f"Excel 解压后不能超过 {MAX_EXCEL_UNCOMPRESSED_BYTES // 1024 // 1024} MB")
        file.stream.seek(0)
        wb = load_workbook(file.stream, read_only=True, data_only=True)
    except ValueError:
        raise
    except zipfile.BadZipFile as exc:
        raise ValueError("Excel 文件损坏或不是有效的 .xlsx 文件") from exc
    except Exception as exc:
        raise ValueError(f"无法读取 Excel：{exc}")
    candidates = []
    for sheet_index, ws in enumerate(wb.worksheets):
        header_rows = list(islice(ws.iter_rows(values_only=True), 8))
        header_idx, mapping = find_import_header(header_rows)
        if header_idx is not None:
            title_hint = 1 if "学生" in ws.title else 0
            candidates.append((title_hint, len(mapping), -sheet_index,
                               ws.title, header_idx, mapping))
    if not candidates:
        wb.close()
        raise ValueError("未识别到学生表头，请至少提供姓名和班级，或使用：年级、班级、学号、姓名、密码")
    _, _, _, sheet_name, header_idx, mapping = max(
        candidates, key=lambda item: (item[0], item[1], item[2]))
    items = []
    class_positions = {}
    normalized_class_rows = 0
    generated_student_nos = 0
    ws = wb[sheet_name]
    try:
        for excel_row_no, row in enumerate(
                ws.iter_rows(min_row=header_idx + 2, values_only=True), header_idx + 2):
            if not any(clean_cell(v) for v in row):
                continue
            if len(items) >= MAX_EXCEL_ROWS:
                raise ValueError(f"一次最多导入 {MAX_EXCEL_ROWS} 名学生，请拆分 Excel 后重试")
            item = {"rowNo": excel_row_no}
            for key in ("grade", "klass", "studentNo", "admissionYear", "classStudentNo",
                        "sourceSequence", "display_name", "username", "password"):
                col = mapping.get(key)
                item[key] = clean_cell(row[col]) if col is not None and col < len(row) else ""
            raw_grade, raw_klass = item["grade"], item["klass"]
            grade, klass, grade_number, class_number, normalized_from, parse_errors = (
                normalize_import_class(raw_grade, raw_klass))
            item["grade"], item["klass"] = grade, klass
            item["classNormalizedFrom"] = normalized_from
            item["parseErrors"] = parse_errors
            if normalized_from:
                normalized_class_rows += 1

            class_key = (grade, klass) if grade and klass else None
            if class_key:
                class_positions[class_key] = class_positions.get(class_key, 0) + 1
            inferred_class_student_no = class_positions.get(class_key) if class_key else None
            if item["classStudentNo"]:
                explicit_class_student_no = parse_class_number(item["classStudentNo"])
                if explicit_class_student_no is None:
                    item["parseErrors"].append("班内号格式不正确")
                    class_student_no = None
                else:
                    class_student_no = explicit_class_student_no
            else:
                class_student_no = inferred_class_student_no
            item["classStudentNo"] = str(class_student_no or "")

            admission_year = None
            if item["admissionYear"]:
                if re.fullmatch(r"\d{4}", item["admissionYear"]):
                    admission_year = int(item["admissionYear"])
                    if not 2000 <= admission_year <= 2100:
                        item["parseErrors"].append("入学年份格式不正确")
                        admission_year = None
                else:
                    item["parseErrors"].append("入学年份应为 4 位年份")
            elif grade_number:
                year_in_stage = grade_year_in_stage(grade_number)
                admission_year = academic_year_start - year_in_stage + 1
            item["admissionYear"] = str(admission_year or "")

            if not item["studentNo"]:
                try:
                    item["studentNo"] = generated_student_no(
                        admission_year, grade_number, class_number, class_student_no)
                    item["studentNoGenerated"] = True
                    generated_student_nos += 1
                except ValueError as exc:
                    item["studentNoGenerated"] = False
                    item["parseErrors"].append(str(exc))
            else:
                item["studentNoGenerated"] = False
            item["username"] = item["username"] or item["studentNo"]
            items.append(item)
    finally:
        wb.close()
    if not items:
        raise ValueError("表格中没有可导入的学生数据")
    return {
        "sheetName": sheet_name,
        "headerRow": header_idx + 1,
        "academicYearStart": academic_year_start,
        "generatedStudentNos": generated_student_nos,
        "normalizedClassRows": normalized_class_rows,
        "rows": items,
    }

def ai_review_student_import(preview):
    return {
        "status": "skipped",
        "summary": "仅执行本地规则校验；未向外部 AI 发送学生名单。",
        "warnings": [],
    }


def build_student_import_preview(file, academic_year_start=None):
    parsed = parse_student_import_file(file, academic_year_start)
    users = load_json(USERS_PATH, {"users": []})["users"]
    classes = load_json(CLASSES_PATH, {"classes": []})
    users_by_username = {u["username"]: u for u in users}
    existing_student_nos = {u.get("studentNo"): u for u in users
                            if u.get("role") == "student" and u.get("studentNo")}
    student_no_counts = {}
    username_counts = {}
    for r in parsed["rows"]:
        if r.get("studentNo"):
            student_no_counts[r["studentNo"]] = student_no_counts.get(r["studentNo"], 0) + 1
        if r.get("username"):
            username_counts[r["username"]] = username_counts.get(r["username"], 0) + 1

    class_keys = {(c.get("grade", ""), c.get("klass", "")) for c in classes.get("classes", [])}
    pending_class_keys = set()
    preview_rows = []
    password_requests = []
    for r in parsed["rows"]:
        errors, warnings = list(r.get("parseErrors") or []), []
        grade = r.get("grade", "")
        klass = r.get("klass", "")
        student_no = r.get("studentNo", "")
        username = r.get("username", "") or student_no
        display_name = r.get("display_name", "")
        for key, label in (("grade", "年级"), ("klass", "班级"), ("studentNo", "学号"), ("display_name", "姓名")):
            if not r.get(key):
                errors.append(f"缺少{label}")
        if student_no and not validate_student_no(student_no):
            errors.append("学号格式不正确")
        if username and not validate_username(username):
            errors.append("用户名格式不正确")
        if r.get("password"):
            policy_error = validate_new_password(r.get("password"))
            if policy_error:
                errors.append(policy_error)
        if student_no_counts.get(student_no, 0) > 1:
            errors.append("Excel 内学号重复")
        if username_counts.get(username, 0) > 1:
            errors.append("Excel 内用户名重复")

        existing_by_no = existing_student_nos.get(student_no)
        existing_by_username = users_by_username.get(username)
        if existing_by_username and existing_by_username.get("role") != "student":
            errors.append("用户名已被非学生账号占用")
        if existing_by_no and existing_by_username and existing_by_no.get("username") != existing_by_username.get("username"):
            errors.append("学号和用户名分别指向不同学生")
        if existing_by_no:
            username = existing_by_no["username"]
            if r.get("username") and r.get("username") != username:
                warnings.append(f"学号已存在，将更新原账号 {username}，不会重命名")
            action = "update"
        elif existing_by_username:
            action = "update"
            if existing_by_username.get("studentNo") and existing_by_username.get("studentNo") != student_no:
                errors.append("用户名已绑定其他学号")
        else:
            action = "new"

        class_action = "exists" if (grade, klass) in class_keys else "create"
        if not grade or not klass:
            class_action = ""
        if class_action == "create":
            pending_class_keys.add((grade, klass))
        status = "error" if errors else action
        password_value = r.get("password") or DEFAULT_TEMP_PASSWORD
        password_hash = ""
        if action == "new" or r.get("password"):
            password_requests.append((len(preview_rows), password_value))
        preview_rows.append({
            **{k: v for k, v in r.items() if k != "password"},
            "username": username,
            "passwordProvided": bool(r.get("password")),
            "passwordHash": password_hash,
            "passwordScheme": "pbkdf2_sha256" if password_hash else "",
            "action": action,
            "classAction": class_action,
            "status": status,
            "errors": errors,
            "warnings": warnings,
        })

    if password_requests:
        with ThreadPoolExecutor(max_workers=min(4, len(password_requests))) as pool:
            hashes = list(pool.map(secure_password_hash,
                                   [password for _index, password in password_requests]))
        for (index, _password), password_hash in zip(password_requests, hashes):
            preview_rows[index]["passwordHash"] = password_hash
            preview_rows[index]["passwordScheme"] = "pbkdf2_sha256"

    summary = {
        "total": len(preview_rows),
        "new": sum(1 for r in preview_rows if r["status"] == "new"),
        "update": sum(1 for r in preview_rows if r["status"] == "update"),
        "error": sum(1 for r in preview_rows if r["status"] == "error"),
        "createClasses": len(pending_class_keys),
        "sheetName": parsed["sheetName"],
        "headerRow": parsed["headerRow"],
        "academicYearStart": parsed["academicYearStart"],
        "generatedStudentNos": parsed["generatedStudentNos"],
        "normalizedClassRows": parsed["normalizedClassRows"],
    }
    preview = {
        "token": secrets.token_hex(16),
        "createdAt": time.time(),
        "summary": summary,
        "rows": preview_rows,
        "classesToCreate": [{"grade": g, "klass": k, "name": class_name(g, k)}
                            for g, k in sorted(pending_class_keys, key=import_class_sort_key)],
    }
    preview["aiReview"] = ai_review_student_import(preview)
    return preview

def validate_import_preview_fresh(preview):
    users = load_json(USERS_PATH, {"users": []}).get("users", [])
    users_by_username = {u.get("username"): u for u in users}
    students_by_no = {}
    duplicate_nos = set()
    for user in users:
        student_no = user.get("studentNo") if user.get("role") == "student" else ""
        if not student_no:
            continue
        if student_no in students_by_no:
            duplicate_nos.add(student_no)
        students_by_no[student_no] = user
    errors = []
    for row in preview.get("rows", []):
        student_no = row.get("studentNo", "")
        username = row.get("username", "")
        by_no = students_by_no.get(student_no)
        by_username = users_by_username.get(username)
        if student_no in duplicate_nos:
            errors.append(f"第 {row.get('rowNo')} 行：现有数据中学号重复")
            continue
        if by_username and by_username.get("role") != "student":
            errors.append(f"第 {row.get('rowNo')} 行：用户名已被非学生账号占用")
            continue
        if by_no and by_username and by_no.get("username") != by_username.get("username"):
            errors.append(f"第 {row.get('rowNo')} 行：学号和用户名已指向不同学生")
            continue
        target = by_no or by_username
        current_action = "update" if target else "new"
        current_username = target.get("username") if target else username
        if current_action != row.get("action") or current_username != username:
            errors.append(f"第 {row.get('rowNo')} 行：账号数据已变化")
        elif target and target.get("studentNo") not in ("", None, student_no):
            errors.append(f"第 {row.get('rowNo')} 行：用户名已绑定其他学号")
    return errors
