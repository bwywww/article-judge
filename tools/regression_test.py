"""隔离回归测试：SQLite、认证、权限、队列、生命周期与结果发布。"""

import json
import os
import tempfile
import sys
import io
import datetime
import sqlite3
import threading
import time
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("APP_ENV", "test")

import app as aiapp
from essay_platform.adapters import ai as ai_adapter
from essay_platform.core import secrets as secret_config
from essay_platform.core.state import scheduled_ocr_jobs
from essay_platform.services.directory import class_sort_key


def reset_runtime(root):
    root = Path(root)
    aiapp.db.reset_state_for_tests()
    paths = {
        "DATA_DIR": str(root),
        "UPLOADS_DIR": str(root / "uploads"),
        "STAGING_DIR": str(root / "uploads" / ".staging"),
        "BACKUPS_DIR": str(root / "backups"),
        "IMPORTS_DIR": str(root / "imports"),
        "CONFIG_PATH": str(root / "config.json"),
        "USERS_PATH": str(root / "users.json"),
        "CLASSES_PATH": str(root / "classes.json"),
        "ASSIGNMENTS_PATH": str(root / "assignments.json"),
        "HISTORY_PATH": str(root / "history.json"),
    }
    paths["DB_PATH"] = aiapp.db.configure(str(root))
    for module in tuple(sys.modules.values()):
        if module and (module.__name__ == "app" or module.__name__.startswith("essay_platform")):
            for name, value in paths.items():
                if hasattr(module, name):
                    setattr(module, name, value)
    aiapp.runtime_initialized_dirs.discard(str(root.resolve()))
    Path(aiapp.UPLOADS_DIR).mkdir(parents=True, exist_ok=True)
    Path(aiapp.STAGING_DIR).mkdir(parents=True, exist_ok=True)
    Path(aiapp.BACKUPS_DIR).mkdir(parents=True, exist_ok=True)
    Path(aiapp.IMPORTS_DIR).mkdir(parents=True, exist_ok=True)
    aiapp.sessions.clear()
    aiapp.in_flight.clear()
    aiapp.scheduled_jobs.clear()
    scheduled_ocr_jobs.clear()
    aiapp.student_submission_in_flight.clear()
    aiapp.upgrade_in_flight.clear()
    aiapp.init_users()
    aiapp.migrate_runtime_data()


def post_json(client, path, body=None, token=None):
    headers = {"X-Token": token} if token else {}
    return client.post(path, json=body or {}, headers=headers)


def put_json(client, path, body=None, token=None):
    headers = {"X-Token": token} if token else {}
    return client.put(path, json=body or {}, headers=headers)


def get_json(client, path, token=None):
    headers = {"X-Token": token} if token else {}
    return client.get(path, headers=headers)


def post_file(client, path, field, file_bytes, filename, token=None, fields=None):
    headers = {"X-Token": token} if token else {}
    data = dict(fields or {})
    data[field] = (io.BytesIO(file_bytes), filename)
    return client.post(path, data=data,
                       content_type="multipart/form-data", headers=headers)


def sample_docx_bytes(names=("张三", "李四")):
    document = Document()
    for index, name in enumerate(names, 1):
        document.add_paragraph("一件有趣的事")
        document.add_paragraph(f"班级：三（1）班　姓名：{name}　学号：{index:02d}")
        document.add_paragraph(f"这是{name}用于验证 Word 分页、姓名和学号识别的作文正文，内容足够长。")
        if index < len(names):
            document.add_page_break()
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def login(client, username, password=None, *, complete_password_change=True):
    password = password or aiapp.DEFAULT_TEMP_PASSWORD
    resp = post_json(client, "/api/login", {"username": username, "password": password})
    assert resp.status_code == 200, resp.get_data(as_text=True)
    payload = resp.get_json()
    token = payload["token"]
    if payload.get("mustChangePassword") and complete_password_change:
        assert get_json(client, "/api/me", token).get_json()["mustChangePassword"] is True
        assert get_json(client, "/api/classes", token).status_code == 428
        changed = post_json(client, "/api/account/password", {
            "currentPassword": password,
            "newPassword": f"{username}-Secure123!",
        }, token)
        assert changed.status_code == 200, changed.get_data(as_text=True)
        token = changed.get_json()["token"]
    return token


def fake_call_llm(messages, max_tokens=4096, temperature=0.6):
    return json.dumps({
        "score": 88,
        "level": "良好",
        "summary": "测试批改结果",
        "dimensions": {"立意与内容": 88, "结构与逻辑": 86, "语言表达": 87},
        "strengths": ["内容完整"],
        "weaknesses": ["细节还可加强"],
        "suggestions": ["补充具体事例"],
        "modelSnippet": "示范句",
        "upgrade": "这是按需生成的完整润色提升版作文内容。",
        "studentName": "",
    }, ensure_ascii=False)


def wait_for(predicate, timeout=6):
    deadline = time.time() + timeout
    while time.time() < deadline:
        if predicate():
            return True
        time.sleep(0.03)
    return False


def wait_for_student_job(client, job_id, token, timeout=8):
    deadline = time.time() + timeout
    latest = None
    while time.time() < deadline:
        response = get_json(client, f"/api/student/jobs/{job_id}", token)
        assert response.status_code == 200, response.get_data(as_text=True)
        latest = response.get_json()
        if latest.get("status") in ("done", "pending_review", "published", "error"):
            return latest
        time.sleep(0.03)
    raise AssertionError(f"学生批改任务超时：{latest}")


def make_student_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "学生名单"
    ws.append(["年级", "班级", "学号", "姓名", "密码"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_teacher_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "教师名单"
    ws.append(["姓名", "工号", "学科", "班级"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_compact_student_workbook():
    wb = Workbook()
    teacher_ws = wb.active
    teacher_ws.title = "班级与任课教师"
    teacher_ws.append(["班级", "语文教师"])
    teacher_ws.append(["三(01)", "测试教师"])
    ws = wb.create_sheet("学生名单")
    ws.append(["序号", "班级", "姓名", "性别"])
    ws.append([1, "三(01)", "导入甲", "男"])
    ws.append([2, "三(1)班", "导入乙", "女"])
    ws.append([3, "三（02）", "导入丙", "男"])
    ws.append([4, "四年级1班", "导入丁", "女"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_sqlite_migration_contract():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        aiapp.db.reset_state_for_tests()
        db_path = Path(aiapp.db.configure(str(root)))
        source = {
            "config": {"base_url": "https://api.deepseek.com", "api_key": "", "model": "mock"},
            "users": {"users": [
                aiapp.make_user("admin", "管理员", "admin"),
                aiapp.make_user("s001", "同学", "student", student_no="20260001"),
            ]},
            "classes": {"classes": [{
                "id": "c1", "name": "一年级一班", "grade": "一年级", "klass": "一班",
                "teacherUsernames": [], "studentUsernames": ["s001"],
            }]},
            "assignments": {"assignments": []},
            "history": {"items": []},
        }
        result = aiapp.db.ensure_database(source)
        assert result["created"] is True and db_path.exists()
        conn = sqlite3.connect(db_path)
        try:
            assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
            assert conn.execute("PRAGMA user_version").fetchone()[0] == aiapp.db.DB_SCHEMA_VERSION
            tables = {row[0] for row in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'")}
            assert {"users", "classes", "assignments", "essays", "history",
                    "feature_modules", "ai_task_routes", "ai_usage"} <= tables
            assert "documents" not in tables
        finally:
            conn.close()
        assert aiapp.db.ensure_database({})["created"] is False

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        sentinel = root / "users.json"
        sentinel.write_text('{"original":true}', encoding="utf-8")
        aiapp.db.reset_state_for_tests()
        db_path = Path(aiapp.db.configure(str(root)))
        duplicate_source = {
            "users": {"users": [
                aiapp.make_user("admin", "管理员", "admin"),
                aiapp.make_user("s001", "甲", "student", student_no="same"),
                aiapp.make_user("s002", "乙", "student", student_no="same"),
            ]},
            "classes": {"classes": []},
            "assignments": {"assignments": []},
            "history": {"items": []},
            "config": {},
        }
        try:
            aiapp.db.ensure_database(duplicate_source)
            raise AssertionError("重复学号迁移必须失败")
        except sqlite3.IntegrityError:
            pass
        assert not db_path.exists()
        assert not Path(str(db_path) + ".migrating").exists()
        assert sentinel.read_text(encoding="utf-8") == '{"original":true}'
    aiapp.db.reset_state_for_tests()


def test_secret_loading_contract():
    names = ("APP_ENV", "DEEPSEEK_KEYS_FILE", "ARTICLE_JUDGE_API_KEY",
             "CODE_DEVELOP_API_KEY", "AI_API_KEY")
    previous = {name: os.environ.get(name) for name in names}
    try:
        with tempfile.TemporaryDirectory() as tmp:
            key_file = Path(tmp) / "deepseekv4.txt"
            key_file.write_text(
                "article_judge sk-article-judge-test-secret\n"
                "code_develop sk-code-develop-test-secret\n",
                encoding="utf-8",
            )
            os.environ["DEEPSEEK_KEYS_FILE"] = str(key_file)
            os.environ["APP_ENV"] = "development"
            for name in ("ARTICLE_JUDGE_API_KEY", "CODE_DEVELOP_API_KEY", "AI_API_KEY"):
                os.environ.pop(name, None)
            loaded = secret_config.load_ai_secrets()
            assert loaded["article_judge"]["source"] == "file"
            assert loaded["code_develop"]["source"] == "file"

            os.environ["CODE_DEVELOP_API_KEY"] = "sk-code-env-test-secret-value"
            assert secret_config.load_ai_secrets()["code_develop"]["source"] == "environment"

            os.environ["APP_ENV"] = "test"
            loaded = secret_config.load_ai_secrets()
            assert loaded["article_judge"]["value"] == ""
            assert loaded["code_develop"]["source"] == "environment"

            os.environ["APP_ENV"] = "production"
            try:
                secret_config.load_ai_secrets()
                raise AssertionError("生产环境不应允许明文密钥文件")
            except secret_config.SecretConfigError:
                pass
    finally:
        for name, value in previous.items():
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value


def test_schema_upgrade_backup():
    with tempfile.TemporaryDirectory() as tmp:
        aiapp.db.reset_state_for_tests()
        db_path = Path(aiapp.db.configure(tmp))
        aiapp.db.ensure_database({
            "users": {"users": []}, "classes": {"classes": []},
            "assignments": {"assignments": []}, "history": {"items": []},
            "config": {},
        })
        with aiapp.db.connection(write=True) as conn:
            conn.execute("""
                INSERT INTO ai_task_routes(task_type,profile,model,updated_at)
                VALUES('development','code_develop','test-model','old')
            """)
            conn.execute("PRAGMA user_version=4")
        aiapp.db.ensure_database({})
        conn = sqlite3.connect(db_path)
        try:
            assert conn.execute("PRAGMA user_version").fetchone()[0] == aiapp.db.DB_SCHEMA_VERSION
            assert conn.execute(
                "SELECT COUNT(*) FROM ai_task_routes WHERE profile='code_develop'"
            ).fetchone()[0] == 0
        finally:
            conn.close()
        backups = list((Path(tmp) / "backups").glob(
            f"schema_v4_to_v{aiapp.db.DB_SCHEMA_VERSION}_*/app.db"))
        assert len(backups) == 1


def main():
    ordered_classes = sorted([
        {"grade": "五年级", "klass": "10班", "name": "五年级10班"},
        {"grade": "高一", "klass": "2班", "name": "高一2班"},
        {"grade": "三年级", "klass": "2班", "name": "三年级2班"},
        {"grade": "五年级", "klass": "1班", "name": "五年级1班"},
        {"grade": "七年级", "klass": "1班", "name": "七年级1班"},
    ], key=class_sort_key)
    assert [item["name"] for item in ordered_classes] == [
        "三年级2班", "五年级1班", "五年级10班", "七年级1班", "高一2班",
    ]
    original_call_llm = aiapp.call_llm
    original_retry_delay = aiapp.GRADE_RETRY_DELAY
    test_sqlite_migration_contract()
    test_secret_loading_contract()
    test_schema_upgrade_backup()
    assert aiapp.default_academic_year_start(datetime.date(2026, 8, 9)) == 2025
    assert aiapp.default_academic_year_start(datetime.date(2026, 9, 1)) == 2026
    assert aiapp.normalize_import_class("", "三(1)班")[:4] == ("三年级", "1班", 3, 1)
    assert aiapp.normalize_import_class("", "三（01）")[:4] == ("三年级", "1班", 3, 1)
    assert aiapp.normalize_import_class("", "四年级1班")[:4] == ("四年级", "1班", 4, 1)
    assert aiapp.grade_year_in_stage(7) == 1
    assert aiapp.grade_year_in_stage(11) == 2
    roster = [
        {"username": "2022040103", "studentNo": "2022040103", "display_name": "测试学生甲"},
        {"username": "2022040114", "studentNo": "2022040114", "display_name": "测试学生乙"},
    ]
    assert aiapp.match_candidate_from_students(
        {"studentName": "测试学生乙", "studentNoHint": "03"}, roster) == "2022040114"
    assert aiapp.match_candidate_from_students(
        {"studentName": "未在名单", "studentNoHint": "03"}, roster) == ""
    assert aiapp.match_candidate_from_students(
        {"studentName": "测试学生乙", "studentNoHint": "2022040103"}, roster) == "2022040114"
    assert aiapp.match_candidate_from_students(
        {"studentName": "", "studentNoHint": "2022040103"}, roster) == "2022040103"
    with tempfile.TemporaryDirectory() as tmp:
        reset_runtime(tmp)
        os.environ["ARTICLE_JUDGE_API_KEY"] = "sk-article-judge-test-secret"
        ai_adapter.call_llm = fake_call_llm
        aiapp.GRADE_RETRY_DELAY = 0.01
        client = aiapp.app.test_client()

        os.environ["ADMIN_PASSWORD"] = "server-secret-123"
        aiapp.apply_admin_password_override()
        users = aiapp.load_json(aiapp.USERS_PATH, {"users": []})["users"]
        system_admin = next(user for user in users if user["username"] == "admin")
        assert aiapp.verify_user_password(system_admin, "server-secret-123")
        assert system_admin["passwordScheme"] == "pbkdf2_sha256"
        system_admin["salt"], system_admin["password_hash"] = aiapp.hash_password("123456")
        system_admin["passwordScheme"] = "legacy_sha256"
        system_admin["mustChangePassword"] = True
        aiapp.save_json(aiapp.USERS_PATH, {"schemaVersion": aiapp.DATA_SCHEMA_VERSION, "users": users})
        del os.environ["ADMIN_PASSWORD"]
        raw_before_migration = {
            path: Path(path).read_bytes()
            for path in (aiapp.USERS_PATH, aiapp.CLASSES_PATH,
                         aiapp.ASSIGNMENTS_PATH, aiapp.HISTORY_PATH)
        }

        admin = login(client, "admin")
        assert Path(aiapp.DB_PATH).exists()
        assert all(Path(path).read_bytes() == content
                   for path, content in raw_before_migration.items())
        assert any(Path(aiapp.BACKUPS_DIR).glob("sqlite_source_*"))
        teacher = login(client, "mr")
        student = login(client, "stu")
        aiapp.sessions.clear()
        assert get_json(client, "/api/me", admin).status_code == 200
        for username in ("admin", "mr", "stu"):
            stored_user = aiapp.db.get_user(username)
            assert stored_user["passwordScheme"] == "pbkdf2_sha256"
            assert stored_user["password_hash"].startswith("pbkdf2:sha256:600000$")
            assert stored_user["mustChangePassword"] is False

        modules_response = get_json(client, "/api/admin/modules", admin)
        assert modules_response.status_code == 200
        modules = {item["key"]: item for item in modules_response.get_json()["modules"]}
        assert modules["auth"]["toggleable"] is False
        assert modules["reporting"]["status"] == "enabled"
        aiapp.db.set_feature_module_state("reporting", "draining", "stale-worker")
        aiapp.db.set_feature_module_state("reporting", "enabled", "admin")
        stale_transition, stale_changed = aiapp.db.compare_and_set_feature_module_state(
            "reporting", "draining", "disabled", "stale-worker")
        assert stale_changed is False and stale_transition["status"] == "enabled"
        assert get_json(client, "/api/admin/modules", teacher).status_code == 403
        assert post_json(client, "/api/admin/modules/auth/state", {
            "state": "disabled",
        }, admin).status_code == 400
        disabled_reporting = post_json(client, "/api/admin/modules/reporting/state", {
            "state": "disabled",
        }, admin)
        assert disabled_reporting.status_code == 200
        assert disabled_reporting.get_json()["module"]["status"] == "disabled"
        assert get_json(client, "/api/student/history", student).status_code == 503
        features = get_json(client, "/api/features", student).get_json()["features"]
        assert features["reporting"] == {"status": "disabled", "available": False}
        assert post_json(client, "/api/admin/modules/reporting/state", {
            "state": "enabled",
        }, admin).status_code == 200

        assert post_json(client, "/api/admin/modules/practice/state", {
            "state": "disabled",
        }, admin).status_code == 200
        practice_disabled = post_json(client, "/api/student/grade", {
            "title": "自主练习", "text": "这是一篇长度足够的自主练习测试作文。",
        }, student)
        assert practice_disabled.status_code == 503
        assert practice_disabled.get_json()["module"] == "practice"
        assert post_json(client, "/api/admin/modules/practice/state", {
            "state": "enabled",
        }, admin).status_code == 200

        assert post_json(client, "/api/admin/modules/imports/state", {
            "state": "disabled",
        }, admin).status_code == 200
        parse_while_disabled = post_file(
            client, "/api/student/parse", "file",
            sample_docx_bytes(("张三",)), "disabled.docx", student,
        )
        assert parse_while_disabled.status_code == 503
        assert parse_while_disabled.get_json()["module"] == "imports"
        assert post_json(client, "/api/admin/modules/imports/state", {
            "state": "enabled",
        }, admin).status_code == 200

        assert post_json(client, "/api/admin/modules/assignments/state", {
            "state": "disabled",
        }, admin).status_code == 200
        cascaded = {
            item["key"]: item["status"] for item in get_json(
                client, "/api/admin/modules", admin).get_json()["modules"]
        }
        assert cascaded["assignments"] == "disabled"
        assert all(cascaded[key] == "disabled"
                   for key in ("imports", "grading", "review", "practice"))
        for module_key in ("assignments", "grading", "practice", "imports", "review"):
            enabled = post_json(client, f"/api/admin/modules/{module_key}/state", {
                "state": "enabled",
            }, admin)
            assert enabled.status_code == 200, enabled.get_data(as_text=True)

        ai_routes = get_json(client, "/api/admin/ai/routes", admin)
        assert ai_routes.status_code == 200
        route_map = {item["taskType"]: item for item in ai_routes.get_json()["routes"]}
        assert route_map["grading"]["model"] == "deepseek-v4-flash"
        assert "development" not in route_map
        assert put_json(client, "/api/admin/ai/routes", {
            "routes": [{"taskType": "development", "model": "deepseek-v4-flash"}],
        }, admin).status_code == 400
        assert get_json(client, "/api/admin/ai/routes", teacher).status_code == 403
        aiapp.db.write_ai_usage(
            profile="article_judge", task_type="grading", model="deepseek-v4-flash",
            actor_username="stu", prompt_tokens=120, completion_tokens=80,
            duration_seconds=1.25, status="ok",
        )
        usage = get_json(client, "/api/admin/ai/usage?page=1&pageSize=20", admin)
        assert usage.status_code == 200
        assert usage.get_json()["summary"]["requests"] >= 1
        assert "text" not in usage.get_json()["items"][0]

        extra_admin_session = login(client, "admin", "admin-Secure123!",
                                    complete_password_change=False)
        assert post_json(client, "/api/logout", {}, extra_admin_session).status_code == 200
        assert get_json(client, "/api/me", extra_admin_session).status_code == 401
        expired_admin_session = login(client, "admin", "admin-Secure123!",
                                      complete_password_change=False)
        expired_hash = aiapp.hashlib.sha256(expired_admin_session.encode("utf-8")).hexdigest()
        with aiapp.db.connection(write=True) as conn:
            conn.execute("UPDATE sessions SET expires_at=0 WHERE token_hash=?", (expired_hash,))
        assert get_json(client, "/api/me", expired_admin_session).status_code == 401

        for attempt in range(aiapp.LOGIN_FAILURE_LIMIT):
            failed = post_json(client, "/api/login", {
                "username": "rate-limit-user", "password": "incorrect-password",
            })
            assert failed.status_code == (429 if attempt == aiapp.LOGIN_FAILURE_LIMIT - 1 else 401)
        blocked = post_json(client, "/api/login", {
            "username": "rate-limit-user", "password": "incorrect-password",
        })
        assert blocked.status_code == 429 and blocked.get_json()["retryAfter"] > 0

        corrupt_path = Path(tmp) / "corrupt.json"
        corrupt_path.write_text("{broken", encoding="utf-8")
        try:
            aiapp.load_json(str(corrupt_path), {})
            raise AssertionError("损坏的 JSON 不应被当成空数据")
        except aiapp.DataFileError:
            pass

        assert post_json(client, "/api/config", {"base_url": "http://mock"}, admin).status_code == 400
        assert post_json(client, "/api/config", {"base_url": "https://api.deepseek.com", "api_key": "k", "model": "m"}, admin).status_code == 200
        os.environ["AI_API_KEY"] = "sk-environment-only-secret-for-tests"
        assert post_json(client, "/api/config", {"api_key": "should-not-be-saved"}, admin).status_code == 200
        managed_config = get_json(client, "/api/config").get_json()
        assert managed_config["managedExternally"] is False and managed_config["has_key"] is True
        assert managed_config["profiles"] == {}
        admin_config = get_json(client, "/api/config", admin).get_json()
        assert admin_config["managedExternally"] is True
        assert set(admin_config["profiles"]) == {"article_judge", "ecnu_plus"}
        web_ai_routes = get_json(client, "/api/admin/ai/routes", admin)
        assert web_ai_routes.status_code == 200
        assert "code_develop" not in web_ai_routes.get_data(as_text=True)
        assert all(route["taskType"] != "development"
                   for route in web_ai_routes.get_json()["routes"])
        assert "code_develop" not in (ROOT / "static" / "js" / "admin.js").read_text(encoding="utf-8")
        assert "code_develop" not in (ROOT / "static" / "js" / "common.js").read_text(encoding="utf-8")
        with aiapp.db.connection() as conn:
            stored_key = conn.execute("SELECT value FROM app_config WHERE key='api_key'").fetchone()[0]
        assert stored_key == ""
        del os.environ["AI_API_KEY"]
        assert post_json(client, "/api/config", {"base_url": "http://mock"}, teacher).status_code == 403
        assert post_json(client, "/api/config/test", {"base_url": "http://mock"}, student).status_code == 403
        single_word = post_file(client, "/api/student/parse", "file",
                                sample_docx_bytes(("张三",)), "single.docx", student)
        assert single_word.status_code == 200, single_word.get_data(as_text=True)
        assert single_word.get_json()["text"].startswith("一件有趣的事")
        multi_word = post_file(client, "/api/student/parse", "file",
                               sample_docx_bytes(), "class.docx", student)
        assert multi_word.status_code == 400
        assert "识别到 2 篇作文" in multi_word.get_json()["error"]
        assert put_json(client, "/api/admin/users/admin", {"role": "student"}, admin).status_code == 400
        assert client.delete("/api/admin/users/admin", headers={"X-Token": admin}).status_code == 400
        assert post_json(client, "/api/admin/users", {
            "role": "admin", "username": "admin2", "display_name": "第二管理员",
        }, admin).status_code == 400
        assert put_json(client, "/api/admin/users/mr", {"display_name": "MR老师2"}, admin).status_code == 200
        assert post_json(client, "/api/student/grade", {
            "title": "老师不能走学生提交",
            "text": "这是一篇用于验证权限的作文内容。",
        }, teacher).status_code == 403
        assert get_json(client, "/api/student/history", admin).status_code == 403

        resp = post_json(client, "/api/admin/users", {
            "role": "student",
            "studentNo": "S010",
            "display_name": "同名学生",
            "password": "S010-Init123!",
        }, admin)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        s010_token = login(client, "S010", "S010-Init123!")
        assert put_json(client, "/api/admin/users/S010", {
            "password": "AnotherPassword123!",
        }, admin).status_code == 400
        reset = post_json(client, "/api/admin/users/S010/reset-password", {}, admin)
        assert reset.status_code == 200
        assert reset.get_json()["temporaryPassword"] == aiapp.DEFAULT_TEMP_PASSWORD
        assert get_json(client, "/api/me", s010_token).status_code == 401
        other_student = login(client, "S010", aiapp.DEFAULT_TEMP_PASSWORD)
        resp = post_json(client, "/api/admin/users", {
            "role": "student",
            "username": "another",
            "studentNo": "S010",
            "display_name": "同名学生",
        }, admin)
        assert resp.status_code == 400

        assert post_json(client, "/api/admin/users", {
            "role": "student",
            "studentNo": "S011",
            "display_name": "同名学生",
        }, admin).status_code == 200
        assert post_json(client, "/api/classes/class/members", {"username": "S010"}, admin).status_code == 200
        assert post_json(client, "/api/classes/class/members", {"username": "S011"}, admin).status_code == 200
        assert aiapp.match_student_username("S010", "class") == "S010"
        assert aiapp.match_student_username("同名学生", "class") == ""

        resp = post_json(client, "/api/classes", {
            "name": "六年级2班", "grade": "六年级", "klass": "2班",
        }, admin)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        other_class_id = resp.get_json()["id"]
        assert post_json(client, f"/api/classes/{other_class_id}/teachers", {
            "username": "mr",
        }, admin).status_code == 200
        assert post_json(client, f"/api/classes/{other_class_id}/members", {
            "username": "S010",
        }, admin).status_code == 200
        classes_data = aiapp.load_json(aiapp.CLASSES_PATH, {"classes": []})["classes"]
        assert sum("S010" in c.get("studentUsernames", []) for c in classes_data) == 1
        assert next(c for c in classes_data if "S010" in c.get("studentUsernames", []))["id"] == other_class_id
        assert post_json(client, "/api/classes/class/members", {"username": "S010"}, admin).status_code == 200
        classes_data = aiapp.load_json(aiapp.CLASSES_PATH, {"classes": []})["classes"]
        assert sum("S010" in c.get("studentUsernames", []) for c in classes_data) == 1
        assert next(c for c in classes_data if "S010" in c.get("studentUsernames", []))["id"] == "class"
        classes_response = get_json(client, "/api/classes", admin).get_json()["classes"]
        assert all("students" not in cls and "teachers" not in cls for cls in classes_response)
        members_response = get_json(client, "/api/classes/class/members?role=student&page=1&pageSize=2", admin)
        assert members_response.status_code == 200
        assert members_response.get_json()["total"] >= 3
        assert len(members_response.get_json()["members"]) == 2
        paged_users = get_json(client, "/api/admin/users?role=student&q=S010&page=1&pageSize=1", admin)
        assert paged_users.status_code == 200
        assert paged_users.get_json()["total"] == 1
        assert paged_users.get_json()["users"][0]["studentNo"] == "S010"
        assert put_json(client, "/api/admin/users/S010", {"grade": "伪造年级"}, admin).status_code == 400

        workbook = make_student_workbook([
            ["五年级", "1班", "S020", "小明", ""],
            ["五年级", "1班", "S021", "小明", "S021-Init123!"],
        ])
        resp = post_file(client, "/api/admin/students/import/preview", "file", workbook, "students.xlsx", admin)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        preview = resp.get_json()
        assert preview["summary"]["new"] == 2
        assert preview["summary"]["error"] == 0
        assert preview["summary"]["createClasses"] == 1
        assert all("password" not in r for r in preview["rows"])
        with aiapp.db.connection() as conn:
            stored_preview = conn.execute(
                "SELECT payload_json FROM import_previews WHERE token=?", (preview["token"],)
            ).fetchone()[0]
        assert "S021-Init123!" not in stored_preview
        assert not any(u.get("studentNo") == "S020" for u in aiapp.load_json(aiapp.USERS_PATH, {"users": []})["users"])
        resp = post_json(client, "/api/admin/students/import/confirm", {"token": preview["token"]}, admin)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        assert post_json(client, "/api/admin/students/import/confirm", {
            "token": preview["token"],
        }, admin).status_code == 409
        users = aiapp.load_json(aiapp.USERS_PATH, {"users": []})["users"]
        assert any(u.get("studentNo") == "S020" and u.get("username") == "S020" for u in users)
        assert login(client, "S020", aiapp.DEFAULT_TEMP_PASSWORD)
        assert login(client, "S021", "S021-Init123!")

        compact_workbook = make_compact_student_workbook()
        resp = post_file(
            client, "/api/admin/students/import/preview", "file", compact_workbook,
            "compact-students.xlsx", admin, fields={"academicYearStart": "2025"})
        assert resp.status_code == 200, resp.get_data(as_text=True)
        compact_preview = resp.get_json()
        compact_rows = compact_preview["rows"]
        assert compact_preview["summary"]["sheetName"] == "学生名单"
        assert compact_preview["summary"]["academicYearStart"] == 2025
        assert compact_preview["summary"]["generatedStudentNos"] == 4
        assert compact_preview["summary"]["normalizedClassRows"] == 4
        assert compact_preview["summary"]["createClasses"] == 3
        assert compact_preview["summary"]["error"] == 0
        assert [item["name"] for item in compact_preview["classesToCreate"]] == [
            "三年级1班", "三年级2班", "四年级1班"]
        assert [row["studentNo"] for row in compact_rows] == [
            "2023030101", "2023030102", "2023030201", "2022040101"]
        assert [(row["grade"], row["klass"]) for row in compact_rows] == [
            ("三年级", "1班"), ("三年级", "1班"),
            ("三年级", "2班"), ("四年级", "1班")]
        assert all(row["studentNoGenerated"] for row in compact_rows)
        original_secure_password_hash = aiapp.secure_password_hash

        def reject_repeated_import_hash(_password):
            raise AssertionError("确认导入不应再次计算学生密码哈希")

        aiapp.secure_password_hash = reject_repeated_import_hash
        confirm_started = time.perf_counter()
        try:
            resp = post_json(client, "/api/admin/students/import/confirm", {
                "token": compact_preview["token"],
            }, admin)
        finally:
            aiapp.secure_password_hash = original_secure_password_hash
        confirm_elapsed = time.perf_counter() - confirm_started
        assert resp.status_code == 200, resp.get_data(as_text=True)
        compact_result = resp.get_json()
        assert compact_result["created"] == 4
        assert compact_result["durationSeconds"] < 5
        assert confirm_elapsed < 5
        users = aiapp.load_json(aiapp.USERS_PATH, {"users": []})["users"]
        assert {u.get("studentNo") for u in users} >= {
            "2023030101", "2023030102", "2023030201", "2022040101"}
        classes_data = aiapp.load_json(aiapp.CLASSES_PATH, {"classes": []})["classes"]
        imported_classes = {(c.get("grade"), c.get("klass")) for c in classes_data}
        assert {("三年级", "1班"), ("三年级", "2班"), ("四年级", "1班")} <= imported_classes

        ai_adapter.call_llm = lambda messages, max_tokens=4096, temperature=0.6: (
            (_ for _ in ()).throw(AssertionError("名单预览不应调用 AI")))
        local_review = aiapp.ai_review_student_import({"summary": {"error": 0}, "rows": []})
        assert local_review["status"] == "skipped"
        assert "本地" in local_review["summary"]
        ai_adapter.call_llm = fake_call_llm

        bad_workbook = make_student_workbook([
            ["六年级", "2班", "S030", "甲", ""],
            ["六年级", "2班", "S030", "乙", ""],
        ])
        resp = post_file(client, "/api/admin/students/import/preview", "file", bad_workbook, "bad.xlsx", admin)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        assert resp.get_json()["summary"]["error"] == 2

        teacher_workbook = make_teacher_workbook([
            ["王老师", "T001", "语文", "class、六年级2班"],
            ["李老师", "T002", "数学", "class"],
        ])
        teacher_preview_response = post_file(
            client, "/api/admin/teachers/import/preview", "file",
            teacher_workbook, "teachers.xlsx", admin)
        assert teacher_preview_response.status_code == 200, teacher_preview_response.get_data(as_text=True)
        teacher_preview = teacher_preview_response.get_json()
        assert teacher_preview["summary"] == {
            "total": 2, "new": 2, "update": 0, "error": 0,
            "sheetName": "教师名单", "headerRow": 1,
        }
        assert not any(user.get("employeeNo") == "T001" for user in aiapp.load_json(
            aiapp.USERS_PATH, {"users": []})["users"])
        teacher_confirm = post_json(
            client, "/api/admin/teachers/import/confirm",
            {"token": teacher_preview["token"]}, admin)
        assert teacher_confirm.status_code == 200, teacher_confirm.get_data(as_text=True)
        assert teacher_confirm.get_json()["created"] == 2
        assert post_json(client, "/api/admin/teachers/import/confirm", {
            "token": teacher_preview["token"],
        }, admin).status_code == 409
        teacher_list = get_json(
            client, "/api/admin/users?role=teacher&q=T001&page=1&pageSize=10", admin).get_json()
        imported_teacher = teacher_list["users"][0]
        assert imported_teacher["employeeNo"] == "T001"
        assert imported_teacher["subject"] == "语文"
        assert {item["id"] for item in imported_teacher["classes"]} == {"class", other_class_id}
        login(client, "T001", aiapp.DEFAULT_TEMP_PASSWORD)

        teacher_update_workbook = make_teacher_workbook([
            ["王老师（更新）", "T001", "语文教研", "class"],
        ])
        update_preview_response = post_file(
            client, "/api/admin/teachers/import/preview", "file",
            teacher_update_workbook, "teachers-update.xlsx", admin)
        assert update_preview_response.status_code == 200
        update_preview = update_preview_response.get_json()
        assert update_preview["summary"]["update"] == 1
        update_confirm = post_json(client, "/api/admin/teachers/import/confirm", {
            "token": update_preview["token"],
        }, admin)
        assert update_confirm.status_code == 200, update_confirm.get_data(as_text=True)
        updated_teacher = get_json(
            client, "/api/admin/users?role=teacher&q=T001&page=1&pageSize=10", admin
        ).get_json()["users"][0]
        assert updated_teacher["display_name"] == "王老师（更新）"
        assert updated_teacher["subject"] == "语文教研"
        assert {item["id"] for item in updated_teacher["classes"]} == {"class"}
        assert login(client, "T001", "T001-Secure123!", complete_password_change=False)

        invalid_teacher_workbook = make_teacher_workbook([
            ["重复甲", "T100", "语文", "class"],
            ["重复乙", "T100", "数学", "不存在班级"],
        ])
        invalid_teacher_preview = post_file(
            client, "/api/admin/teachers/import/preview", "file",
            invalid_teacher_workbook, "teachers-invalid.xlsx", admin)
        assert invalid_teacher_preview.status_code == 200
        assert invalid_teacher_preview.get_json()["summary"]["error"] == 2
        assert post_json(client, "/api/admin/teachers/import/confirm", {
            "token": invalid_teacher_preview.get_json()["token"],
        }, admin).status_code == 400

        disposable_teacher = post_json(client, "/api/admin/users", {
            "role": "teacher", "employeeNo": "TDEL", "display_name": "可删除老师",
            "subject": "美术", "classIds": ["class"],
        }, admin)
        assert disposable_teacher.status_code == 200, disposable_teacher.get_data(as_text=True)
        removed_disposable = client.delete(
            "/api/admin/users/TDEL", headers={"X-Token": admin})
        assert removed_disposable.status_code == 200
        assert removed_disposable.get_json()["deleted"] == ["TDEL"]
        assert aiapp.db.get_user("TDEL") is None

        business_teacher = post_json(client, "/api/admin/users", {
            "role": "teacher", "employeeNo": "TBIZ", "display_name": "业务老师",
            "subject": "语文", "classIds": ["class", other_class_id],
        }, admin)
        assert business_teacher.status_code == 200
        business_token = login(client, "TBIZ", aiapp.DEFAULT_TEMP_PASSWORD)
        business_assignment = post_json(client, "/api/assignments", {
            "title": "待移交作业", "classId": "class",
        }, business_token)
        assert business_assignment.status_code == 200
        business_assignment_id = business_assignment.get_json()["id"]
        deactivated_business = client.delete(
            "/api/admin/users/TBIZ", headers={"X-Token": admin})
        assert deactivated_business.status_code == 200
        assert deactivated_business.get_json()["deactivated"] == ["TBIZ"]
        assert get_json(client, "/api/me", business_token).status_code == 401
        assert aiapp.db.get_user("TBIZ")["active"] is False
        reenabled = put_json(client, "/api/admin/users/TBIZ", {
            "employeeNo": "TBIZ2", "display_name": "业务老师（复职）", "subject": "语文",
            "classIds": ["class"], "active": True,
        }, admin)
        assert reenabled.status_code == 200, reenabled.get_data(as_text=True)
        assert reenabled.get_json()["username"] == "TBIZ2"
        assert aiapp.db.get_user("TBIZ") is None and aiapp.db.get_user("TBIZ2")["active"] is True
        renamed_assignment = aiapp.find_assignment(
            aiapp.load_json(aiapp.ASSIGNMENTS_PATH, {"assignments": []}), business_assignment_id)
        assert renamed_assignment["createdBy"] == "TBIZ2"
        assert login(client, "TBIZ2", "TBIZ-Secure123!", complete_password_change=False)

        batch_removed = post_json(client, "/api/admin/teachers/batch-delete", {
            "usernames": ["T002", "TBIZ2", "missing-teacher"],
        }, admin)
        assert batch_removed.status_code == 200
        batch_payload = batch_removed.get_json()
        assert batch_payload["deleted"] == ["T002"]
        assert batch_payload["deactivated"] == ["TBIZ2"]
        assert batch_payload["failed"] == [{"username": "missing-teacher", "reason": "老师账号不存在"}]
        assert aiapp.find_assignment(
            aiapp.load_json(aiapp.ASSIGNMENTS_PATH, {"assignments": []}), business_assignment_id)

        stale_workbook = make_student_workbook([
            ["六年级", "3班", "future_teacher", "冲突账号", ""],
        ])
        resp = post_file(client, "/api/admin/students/import/preview", "file",
                         stale_workbook, "stale.xlsx", admin)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        stale_token = resp.get_json()["token"]
        assert post_json(client, "/api/admin/users", {
            "role": "teacher", "employeeNo": "future_teacher", "display_name": "新增老师",
            "subject": "语文", "classIds": ["class"],
        }, admin).status_code == 200
        resp = post_json(client, "/api/admin/students/import/confirm", {
            "token": stale_token,
        }, admin)
        assert resp.status_code == 409, resp.get_data(as_text=True)

        audit = get_json(client, "/api/admin/audit?page=1&pageSize=100", admin)
        assert audit.status_code == 200
        audit_payload = audit.get_json()
        audit_types = {item["type"] for item in audit_payload["items"]}
        assert {"login_succeeded", "password_changed", "password_reset",
                "student_import_confirmed", "ai_config_updated"} <= audit_types
        audit_text = json.dumps(audit_payload, ensure_ascii=False)
        assert "environment-only-secret" not in audit_text
        assert "admin-Secure123!" not in audit_text
        users = aiapp.load_json(aiapp.USERS_PATH, {"users": []})["users"]
        assert next(u for u in users if u["username"] == "future_teacher")["role"] == "teacher"

        assert client.get("/student").status_code == 302
        assert client.get("/student").headers["Location"].endswith("/student/assignments")
        assert client.get("/student/assignments").status_code == 200
        assert client.get("/student/practice").status_code == 200
        assert client.get("/student/history").status_code == 200
        assert b"btn-settings" not in client.get("/").data
        assert b"btn-settings" not in client.get("/teacher").data
        assert b"btn-settings" not in client.get("/student/assignments").data
        assert b"btn-settings" not in client.get("/student/practice").data
        assert b"btn-settings" not in client.get("/student/history").data
        admin_page = client.get("/admin").data
        assert b"btn-settings" not in admin_page
        assert b'data-tab="modules"' in admin_page
        assert b'data-tab="ai"' in admin_page

        future_deadline = aiapp.utc_iso(aiapp.utc_now() + aiapp.datetime.timedelta(hours=2))
        resp = post_json(client, "/api/assignments", {
            "title": "生命周期测试", "prompt": "写一篇作文", "classId": "class",
            "deadlineAt": future_deadline,
        }, teacher)
        lifecycle_id = resp.get_json()["id"]
        assert resp.get_json()["lifecycleStatus"] == "draft"
        assert lifecycle_id not in {item["id"] for item in get_json(
            client, "/api/assignments", student).get_json()["assignments"]}
        assert post_json(client, f"/api/assignments/{lifecycle_id}/publish", {}, teacher).status_code == 200
        assert lifecycle_id in {item["id"] for item in get_json(
            client, "/api/assignments", student).get_json()["assignments"]}
        assert post_json(client, f"/api/assignments/{lifecycle_id}/unpublish", {}, teacher).status_code == 200
        assert lifecycle_id not in {item["id"] for item in get_json(
            client, "/api/assignments", student).get_json()["assignments"]}
        assert post_json(client, f"/api/assignments/{lifecycle_id}/publish", {}, teacher).status_code == 200
        assert post_json(client, f"/api/assignments/{lifecycle_id}/essays", {
            "essays": [{"studentName": "生命周期学生", "text": "用于冻结作业内容的老师导入作文正文。"}],
        }, teacher).status_code == 200
        assert put_json(client, f"/api/assignments/{lifecycle_id}", {
            "title": "不应修改",
        }, teacher).status_code == 409
        assert post_json(client, f"/api/assignments/{lifecycle_id}/unpublish", {}, teacher).status_code == 409
        assert post_json(client, f"/api/assignments/{lifecycle_id}/close", {}, teacher).status_code == 200
        closed_item = next(item for item in get_json(
            client, "/api/assignments", student).get_json()["assignments"]
                           if item["id"] == lifecycle_id)
        assert closed_item["lifecycleStatus"] == "closed" and closed_item["canSubmit"] is False
        assert post_json(client, f"/api/assignments/{lifecycle_id}/reopen", {}, teacher).status_code == 200
        assert post_json(client, f"/api/assignments/{lifecycle_id}/close", {}, teacher).status_code == 200
        assert post_json(client, f"/api/assignments/{lifecycle_id}/archive", {}, teacher).status_code == 200
        assert lifecycle_id not in {item["id"] for item in get_json(
            client, "/api/assignments", student).get_json()["assignments"]}
        assert client.delete(f"/api/assignments/{lifecycle_id}", headers={"X-Token": teacher}).status_code == 409

        deletable = post_json(client, "/api/assignments", {
            "title": "可删除空草稿", "classId": "class",
        }, teacher).get_json()["id"]
        assert client.delete(f"/api/assignments/{deletable}", headers={"X-Token": teacher}).status_code == 200

        expiring = post_json(client, "/api/assignments", {
            "title": "到期自动关闭", "classId": "class", "deadlineAt": future_deadline,
        }, teacher).get_json()["id"]
        assert post_json(client, f"/api/assignments/{expiring}/publish", {}, teacher).status_code == 200
        past_deadline = aiapp.utc_iso(aiapp.utc_now() - aiapp.datetime.timedelta(minutes=1))
        with aiapp.db.connection(write=True) as conn:
            conn.execute("UPDATE assignments SET deadline_at=? WHERE id=?", (past_deadline, expiring))
        expiring_item = next(item for item in get_json(
            client, "/api/assignments", student).get_json()["assignments"]
                             if item["id"] == expiring)
        assert expiring_item["lifecycleStatus"] == "closed"
        assert post_json(client, "/api/student/grade", {
            "assignmentId": expiring, "text": "到期后不允许再提交这篇作文。",
        }, student).status_code == 409
        assert post_json(client, f"/api/assignments/{expiring}/reopen", {}, teacher).status_code == 400
        next_deadline = aiapp.utc_iso(aiapp.utc_now() + aiapp.datetime.timedelta(hours=3))
        assert post_json(client, f"/api/assignments/{expiring}/reopen", {
            "deadlineAt": next_deadline,
        }, teacher).status_code == 200

        resp = post_json(client, "/api/assignments", {"title": "锁定测试", "prompt": "写一篇作文", "classId": "class"}, teacher)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        aid = resp.get_json()["id"]
        assert put_json(client, f"/api/assignments/{aid}", {"classId": ""}, teacher).status_code == 400

        resp = post_json(client, f"/api/assignments/{aid}/essays", {
            "fileKey": "../../users.json",
            "essays": [{"studentName": "STU学生", "submittedBy": "stu",
                        "text": "这是路径穿越测试作文内容，不能接受旧版文件标识。"}],
        }, teacher)
        assert resp.status_code == 400
        resp = post_json(client, f"/api/assignments/{aid}/essays", {
            "essays": [
                {"studentName": "学生甲", "submittedBy": "stu", "text": "第一篇重复绑定测试作文内容。"},
                {"studentName": "学生乙", "submittedBy": "stu", "text": "第二篇重复绑定测试作文内容。"},
            ],
        }, teacher)
        assert resp.status_code == 409

        unbound_payload = {
            "essays": [{
                "studentName": "未绑定匿名作文",
                "text": "这是同一来源重复确认时必须幂等更新、不能不断追加副本的测试作文内容。",
            }],
        }
        assert post_json(client, f"/api/assignments/{aid}/essays", unbound_payload, teacher).status_code == 200
        first_unbound = [
            essay for essay in aiapp.find_assignment(
                aiapp.load_json(aiapp.ASSIGNMENTS_PATH, {"assignments": []}), aid
            )["essays"]
            if essay.get("studentName") == "未绑定匿名作文"
        ]
        assert len(first_unbound) == 1
        first_unbound_id = first_unbound[0]["id"]
        assert post_json(client, f"/api/assignments/{aid}/essays", unbound_payload, teacher).status_code == 200
        repeated_unbound = [
            essay for essay in aiapp.find_assignment(
                aiapp.load_json(aiapp.ASSIGNMENTS_PATH, {"assignments": []}), aid
            )["essays"]
            if essay.get("studentName") == "未绑定匿名作文"
        ]
        assert len(repeated_unbound) == 1 and repeated_unbound[0]["id"] == first_unbound_id

        resp = post_json(client, f"/api/assignments/{aid}/essays", {
            "essays": [{"studentName": "STU学生", "submittedBy": "stu", "text": "这是一篇用于老师导入锁定测试的作文内容。"}]
        }, teacher)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        assert put_json(client, f"/api/assignments/{aid}", {
            "classId": other_class_id,
        }, teacher).status_code == 409

        assignments_data = aiapp.load_json(aiapp.ASSIGNMENTS_PATH, {"assignments": []})
        assignment = aiapp.find_assignment(assignments_data, aid)
        assignment["sourceFile"] = "../../users.json"
        aiapp.save_json(aiapp.ASSIGNMENTS_PATH, assignments_data)
        assert get_json(client, f"/api/assignments/{aid}/source", teacher).status_code == 404
        assignment.pop("sourceFile", None)
        aiapp.save_json(aiapp.ASSIGNMENTS_PATH, assignments_data)

        resp = post_file(client, f"/api/assignments/{aid}/upload", "file",
                         sample_docx_bytes(), "class-essays.docx", teacher)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        word_upload = resp.get_json()
        assert [essay["studentName"] for essay in word_upload["essays"]] == ["张三", "李四"]
        assert [essay["studentNoHint"] for essay in word_upload["essays"]] == ["01", "02"]
        assert all("班级：" not in essay["text"] and "姓名：" not in essay["text"]
                   for essay in word_upload["essays"])
        assert post_json(client, f"/api/assignments/{aid}/upload/cancel", {
            "uploadToken": word_upload["uploadToken"],
        }, teacher).status_code == 200

        ai_split_source = (
            "第一篇作文开头内容，后面是第一篇完整且足够长的正文。\n"
            "第二篇作文开头内容，后面是第二篇完整且足够长的正文。"
        )
        def split_call(messages, max_tokens=4096, temperature=0.6):
            if "startMarker" in messages[-1]["content"]:
                return json.dumps({"segments": [
                    {"studentName": "张三", "startMarker": "第一篇作文开头内容"},
                    {"studentName": "李四", "startMarker": "第二篇作文开头内容"},
                ]}, ensure_ascii=False)
            return fake_call_llm(messages, max_tokens, temperature)

        ai_adapter.call_llm = split_call
        resp = post_json(client, f"/api/assignments/{aid}/ai-split", {
            "studentName": "第1篇", "text": ai_split_source,
        }, teacher)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        split_result = resp.get_json()
        assert split_result["split"] is True
        assert [essay["studentName"] for essay in split_result["essays"]] == ["张三", "李四"]
        assert split_result["essays"][0]["text"].startswith("第一篇作文开头内容")
        assert split_result["essays"][1]["text"].startswith("第二篇作文开头内容")
        assert "\n".join(essay["text"] for essay in split_result["essays"]) == ai_split_source
        ai_adapter.call_llm = fake_call_llm

        upload_text = "姓名：张三\n我的家乡\n这是一篇用于验证临时上传原件确认流程的测试作文内容。".encode("utf-8")
        resp = post_file(client, f"/api/assignments/{aid}/upload", "file",
                         upload_text, "essays.txt", teacher)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        uploaded = resp.get_json()
        assert uploaded.get("uploadToken") and not uploaded.get("fileKey")
        uploaded["essays"][0]["submittedBy"] = "stu"
        resp = post_json(client, f"/api/assignments/{aid}/essays", {
            "essays": uploaded["essays"], "uploadToken": uploaded["uploadToken"],
        }, teacher)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        assignment = aiapp.find_assignment(
            aiapp.load_json(aiapp.ASSIGNMENTS_PATH, {"assignments": []}), aid)
        source_file = assignment["sourceFile"]
        source_path = Path(aiapp.UPLOADS_DIR) / aid / source_file
        assert source_path.is_file()
        assert get_json(client, f"/api/assignments/{aid}/source", teacher).status_code == 200

        resp = post_file(client, f"/api/assignments/{aid}/upload", "file",
                         upload_text, "replacement.txt", teacher)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        cancel_token = resp.get_json()["uploadToken"]
        assert source_path.is_file()
        assert post_json(client, f"/api/assignments/{aid}/upload/cancel", {
            "uploadToken": cancel_token,
        }, teacher).status_code == 200
        assert source_path.is_file()
        assert post_json(client, f"/api/assignments/{aid}/publish", {}, teacher).status_code == 200

        assignments = get_json(client, "/api/assignments", student).get_json()["assignments"]
        locked = next(a for a in assignments if a["id"] == aid)
        assert locked["submissionStatus"] == "teacher_submitted_locked"
        assert locked["lockedByTeacher"] is True

        resp = post_json(client, "/api/student/grade", {
            "assignmentId": aid,
            "title": "锁定测试",
            "prompt": "写一篇作文",
            "text": "学生试图覆盖老师导入内容，这里应该被拒绝。",
        }, student)
        assert resp.status_code == 409

        resp = post_json(client, "/api/assignments", {"title": "自提交测试", "prompt": "写一篇作文", "classId": "class"}, teacher)
        aid2 = resp.get_json()["id"]
        assert post_json(client, f"/api/assignments/{aid2}/publish", {}, teacher).status_code == 200
        # 保存后未开始批改的老师导入作文应保持 pending，用于前端显示“开始批改”而不是假装批改中。
        detail = get_json(client, f"/api/assignments/{aid}", teacher).get_json()
        assert detail["essays"][0]["status"] == "pending"
        captured_messages = []
        captured_started = threading.Event()
        captured_release = threading.Event()

        def capture_call(messages, max_tokens=4096, temperature=0.6):
            captured_messages.append(messages)
            captured_started.set()
            captured_release.wait(5)
            return fake_call_llm(messages, max_tokens, temperature)

        ai_adapter.call_llm = capture_call
        submit_started_at = time.time()
        resp = post_json(client, "/api/student/grade", {
            "assignmentId": aid2,
            "title": "客户端伪造题目",
            "prompt": "客户端伪造要求",
            "text": "这是一篇普通学生自提交作文，应该可以保存并显示已提交。",
        }, student)
        assert resp.status_code == 202, resp.get_data(as_text=True)
        assert time.time() - submit_started_at < 1.0
        first_job_id = resp.get_json()["jobId"]
        assert captured_started.wait(3)
        duplicate = post_json(client, "/api/student/grade", {
            "assignmentId": aid2,
            "text": "这是一篇普通学生自提交作文，应该可以保存并显示已提交。",
        }, student)
        assert duplicate.status_code == 202
        assert duplicate.get_json()["jobId"] == first_job_id
        captured_release.set()
        first_job = wait_for_student_job(client, first_job_id, student)
        assert first_job["status"] == "pending_review"
        assert "result" not in first_job
        assert get_json(client, f"/api/student/jobs/{first_job_id}", other_student).status_code == 404
        assert get_json(client, f"/api/student/jobs/{first_job_id}", teacher).status_code == 403
        prompt_text = json.dumps(captured_messages, ensure_ascii=False)
        assert "客户端伪造题目" in prompt_text and "写一篇作文" in prompt_text
        assert "客户端伪造要求" not in prompt_text
        ai_adapter.call_llm = fake_call_llm
        assignments = get_json(client, "/api/assignments", student).get_json()["assignments"]
        submitted = next(a for a in assignments if a["id"] == aid2)
        assert submitted["submissionStatus"] == "pending_review"
        assert submitted["studentTitle"] == "客户端伪造题目"
        teacher_detail = get_json(client, f"/api/assignments/{aid2}", teacher).get_json()
        assert next(item for item in teacher_detail["essays"]
                    if item["submittedBy"] == "stu")["studentTitle"] == "客户端伪造题目"
        assert submitted["submissionLocked"] is True
        history = get_json(client, "/api/student/history", student).get_json()["history"]
        assert not any(i.get("assignmentId") == aid2 for i in history)

        publish_first = post_json(
            client, f"/api/assignments/{aid2}/essays/{first_job['submissionId']}/publish-result",
            {}, teacher)
        assert publish_first.status_code == 200, publish_first.get_data(as_text=True)
        first_history_id = publish_first.get_json()["published"][0]["historyId"]
        published_job = get_json(client, f"/api/student/jobs/{first_job_id}", student).get_json()
        assert published_job["status"] == "published"
        assert published_job["historyId"] == first_history_id
        published_assignment = next(item for item in get_json(
            client, "/api/assignments", student).get_json()["assignments"]
                                    if item["id"] == aid2)
        assert published_assignment["publishedScore"] == 88
        assert published_assignment["resultPublishedAt"]
        assert post_json(
            client, f"/api/assignments/{aid2}/essays/{first_job['submissionId']}/publish-result",
            {}, teacher).status_code == 409
        history = get_json(client, "/api/student/history", student).get_json()["history"]
        assert sum(i.get("assignmentId") == aid2 for i in history) == 1

        review_detail = get_json(client, f"/api/assignments/{aid2}", teacher).get_json()
        review_essay = next(item for item in review_detail["essays"]
                            if item["id"] == first_job["submissionId"])
        original_version = review_essay["feedbackVersion"]
        edited_feedback = {
            **review_essay["feedback"],
            "score": 92,
            "level": "优",
            "summary": "老师人工复核后的总评",
            "dimensions": {"立意与内容": 92, "结构与逻辑": 90, "语言表达": 91},
            "strengths": ["老师确认内容具体"],
            "weaknesses": ["结尾仍可收束"],
            "suggestions": ["补写结尾照应题目"],
            "modelSnippet": "老师修改后的示范句",
            "upgrade": "老师保存的整篇提升版。",
        }
        forbidden_edit = put_json(
            client, f"/api/assignments/{aid2}/essays/{first_job['submissionId']}/feedback",
            {"feedbackVersion": original_version, "feedback": edited_feedback}, admin)
        assert forbidden_edit.status_code == 403
        saved_edit = put_json(
            client, f"/api/assignments/{aid2}/essays/{first_job['submissionId']}/feedback",
            {"feedbackVersion": original_version, "feedback": edited_feedback}, teacher)
        assert saved_edit.status_code == 200, saved_edit.get_data(as_text=True)
        assert saved_edit.get_json()["wasPublished"] is True
        assert saved_edit.get_json()["feedbackVersion"] == original_version + 1
        conflict = put_json(
            client, f"/api/assignments/{aid2}/essays/{first_job['submissionId']}/feedback",
            {"feedbackVersion": original_version, "feedback": edited_feedback}, teacher)
        assert conflict.status_code == 409
        pending_detail = get_json(client, f"/api/assignments/{aid2}", teacher).get_json()
        pending_essay = next(item for item in pending_detail["essays"]
                             if item["id"] == first_job["submissionId"])
        assert pending_essay["resultStatus"] == "pending_review"
        assert pending_essay["score"] == 92 and pending_essay["level"] == "优"
        assert pending_essay["feedbackEditedBy"] == "mr"
        student_pending = next(item for item in get_json(
            client, "/api/assignments", student).get_json()["assignments"] if item["id"] == aid2)
        assert student_pending["publishedScore"] == 88
        old_published_snapshot = next(item for item in get_json(
            client, "/api/student/history", student).get_json()["history"]
                                      if item["id"] == first_history_id)
        assert old_published_snapshot["result"]["score"] == 88
        republished_edit = post_json(
            client, f"/api/assignments/{aid2}/essays/{first_job['submissionId']}/publish-result",
            {}, teacher)
        assert republished_edit.status_code == 200
        edited_history_id = republished_edit.get_json()["published"][0]["historyId"]
        edited_snapshot = next(item for item in get_json(
            client, "/api/student/history", student).get_json()["history"]
                               if item["id"] == edited_history_id)
        assert edited_snapshot["result"]["score"] == 92
        assert edited_snapshot["result"]["summary"] == "老师人工复核后的总评"
        unchanged_edit = put_json(
            client, f"/api/assignments/{aid2}/essays/{first_job['submissionId']}/feedback",
            {"feedbackVersion": original_version + 1, "feedback": edited_feedback}, teacher)
        assert unchanged_edit.status_code == 200
        assert unchanged_edit.get_json()["unchanged"] is True
        unchanged_detail = get_json(client, f"/api/assignments/{aid2}", teacher).get_json()
        unchanged_essay = next(item for item in unchanged_detail["essays"]
                               if item["id"] == first_job["submissionId"])
        assert unchanged_essay["resultStatus"] == "published"
        assert unchanged_essay["feedbackVersion"] == original_version + 1
        with aiapp.db.connection() as conn:
            stored = conn.execute(
                "SELECT feedback_json,ai_feedback_json FROM essays WHERE id=?",
                (first_job["submissionId"],)).fetchone()
        assert json.loads(stored["feedback_json"])["score"] == 92
        assert json.loads(stored["ai_feedback_json"])["score"] == 88

        assert post_json(client, "/api/student/grade", {
            "assignmentId": aid2,
            "text": "没有经过老师重新开放时，学生不能重复覆盖第一次提交。",
        }, student).status_code == 409
        assert post_json(client, f"/api/assignments/{aid2}/submissions/stu/reopen", {}, teacher).status_code == 200
        assignments = get_json(client, "/api/assignments", student).get_json()["assignments"]
        reopened = next(a for a in assignments if a["id"] == aid2)
        assert reopened["submissionStatus"] == "resubmission_open"
        assert reopened["submissionLocked"] is False
        resp = post_json(client, "/api/student/grade", {
            "assignmentId": aid2,
            "title": "仍然忽略的伪造题目",
            "text": "老师已经重新开放，因此这次修改后的作文可以重新提交并再次锁定。",
        }, student)
        assert resp.status_code == 202, resp.get_data(as_text=True)
        second_job = wait_for_student_job(client, resp.get_json()["jobId"], student)
        assert second_job["status"] == "pending_review"
        old_snapshot = next(item for item in get_json(
            client, "/api/student/history", student).get_json()["history"]
                            if item["id"] == first_history_id)
        assert post_json(client, f"/api/assignments/{aid2}/close", {}, teacher).status_code == 200
        assert post_json(client, f"/api/assignments/{aid2}/archive", {}, teacher).status_code == 409
        publish_second = post_json(client, f"/api/assignments/{aid2}/publish-results", {}, teacher)
        assert publish_second.status_code == 200
        assert publish_second.get_json()["count"] == 1
        second_history_id = publish_second.get_json()["published"][0]["historyId"]
        preserved_snapshot = next(item for item in get_json(
            client, "/api/student/history", student).get_json()["history"]
                                  if item["id"] == first_history_id)
        assert preserved_snapshot["result"] == old_snapshot["result"]
        assert post_json(client, f"/api/assignments/{aid2}/reopen", {}, teacher).status_code == 200
        assignments = get_json(client, "/api/assignments", student).get_json()["assignments"]
        assert next(a for a in assignments if a["id"] == aid2)["submissionStatus"] == "result_published"

        resp = post_json(client, f"/api/assignments/{aid2}/essays", {
            "essays": [{"studentName": "STU学生", "submittedBy": "stu",
                        "text": "老师覆盖学生提交前必须得到后端明确确认。"}],
        }, teacher)
        assert resp.status_code == 409
        resp = post_json(client, f"/api/assignments/{aid2}/essays", {
            "overrideStudentSubmissions": True,
            "essays": [{"studentName": "STU学生", "submittedBy": "stu",
                        "text": "老师确认覆盖后，学生端应恢复老师导入锁定状态。"}],
        }, teacher)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        assignments = get_json(client, "/api/assignments", student).get_json()["assignments"]
        assert next(a for a in assignments if a["id"] == aid2)["submissionStatus"] == "teacher_submitted_locked"

        resp = post_json(client, f"/api/student/history/{second_history_id}/upgrade", {}, student)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        assert resp.get_json()["upgrade"] == "这是按需生成的完整润色提升版作文内容。"
        assert post_json(client, f"/api/student/history/{first_history_id}/upgrade", {}, teacher).status_code == 403

        resp = post_json(client, "/api/student/grade", {
            "title": "自主练习",
            "prompt": "自由练习",
            "text": "这是一篇普通自主练习作文，不应该关联任何老师作业。",
        }, student)
        assert resp.status_code == 202, resp.get_data(as_text=True)
        practice_job = wait_for_student_job(client, resp.get_json()["jobId"], student)
        assert practice_job["status"] == "done"
        practice_history_id = practice_job["historyId"]
        assert practice_job["result"].get("upgrade") == ""
        resp = post_json(client, f"/api/student/history/{practice_history_id}/upgrade", {}, student)
        assert resp.status_code == 200, resp.get_data(as_text=True)
        assert resp.get_json()["upgrade"] == "这是按需生成的完整润色提升版作文内容。"
        assert post_json(client, f"/api/student/history/{practice_history_id}/upgrade", {}, teacher).status_code == 403
        history = get_json(client, "/api/student/history", student).get_json()["history"]
        assert sum(i.get("assignmentId") == aid2 for i in history) >= 2
        assert any(not i.get("assignmentId") for i in history)
        assignment_history = get_json(
            client, "/api/student/history?type=assignment&page=1&pageSize=1", student).get_json()
        practice_history = get_json(
            client, "/api/student/history?type=practice&page=1&pageSize=1", student).get_json()
        assert assignment_history["total"] >= 2 and len(assignment_history["history"]) == 1
        assert all(item.get("assignmentId") for item in assignment_history["history"])
        assert practice_history["total"] >= 1 and len(practice_history["history"]) == 1
        assert all(not item.get("assignmentId") for item in practice_history["history"])
        records = get_json(client, "/api/teacher/records", teacher).get_json()["records"]
        assert any(r.get("assignmentId") == aid2 for r in records)
        assert all(r.get("assignmentId") for r in records)
        published_audit = get_json(
            client, "/api/admin/audit?type=assignment_result_published&page=1&pageSize=10", admin)
        assert published_audit.status_code == 200 and published_audit.get_json()["total"] >= 1

        batch_ids = []
        for batch_no in range(2):
            resp = post_json(client, "/api/assignments", {
                "title": f"全局队列测试{batch_no + 1}", "prompt": "统一题目", "classId": "class",
            }, teacher)
            batch_id = resp.get_json()["id"]
            batch_ids.append(batch_id)
            essays = [{
                "studentName": f"队列学生{batch_no}-{idx}",
                "text": f"这是第{idx}篇用于验证全局批改并发上限的作文内容，长度足够。",
            } for idx in range(4)]
            assert post_json(client, f"/api/assignments/{batch_id}/essays", {
                "essays": essays,
            }, teacher).status_code == 200

        counter_lock = threading.Lock()
        release_jobs = threading.Event()
        active = 0
        max_active = 0

        def blocking_call(messages, max_tokens=4096, temperature=0.6):
            nonlocal active, max_active
            with counter_lock:
                active += 1
                max_active = max(max_active, active)
            release_jobs.wait(5)
            with counter_lock:
                active -= 1
            return fake_call_llm(messages, max_tokens, temperature)

        ai_adapter.call_llm = blocking_call
        first_start = post_json(client, f"/api/assignments/{batch_ids[0]}/grade", {}, teacher)
        assert first_start.status_code == 200 and first_start.get_json()["started"] == 4
        assert wait_for(lambda: max_active == aiapp.GRADE_CONCURRENCY)
        duplicate_start = post_json(client, f"/api/assignments/{batch_ids[0]}/grade", {}, teacher)
        assert duplicate_start.status_code == 200
        assert duplicate_start.get_json()["started"] == 0
        assert duplicate_start.get_json()["existing"] == 4
        second_start = post_json(client, f"/api/assignments/{batch_ids[1]}/grade", {}, teacher)
        assert second_start.status_code == 200 and second_start.get_json()["started"] == 4
        shared_student_job = post_json(client, "/api/student/grade", {
            "title": "共享队列自主练习",
            "text": "这篇自主练习必须与老师批改任务共用同一个四路线程池。",
        }, student)
        assert shared_student_job.status_code == 202
        time.sleep(0.12)
        assert max_active == aiapp.GRADE_CONCURRENCY
        queued_shared = aiapp.db.get_grading_job(shared_student_job.get_json()["jobId"])
        assert queued_shared["status"] == "queued"
        draining = post_json(client, "/api/admin/modules/grading/state", {
            "state": "disabled",
        }, admin)
        assert draining.status_code == 200
        assert draining.get_json()["module"]["status"] == "draining"
        blocked_during_drain = post_json(client, "/api/student/grade", {
            "title": "排空期间不应接收", "text": "这个新任务应该被模块状态机拒绝。",
        }, student)
        assert blocked_during_drain.status_code == 503
        assert blocked_during_drain.get_json()["state"] == "draining"
        release_jobs.set()

        def batches_finished():
            data = aiapp.load_json(aiapp.ASSIGNMENTS_PATH, {"assignments": []})
            return all(all(e.get("status") == "done" for e in aiapp.find_assignment(data, batch_id)["essays"])
                       for batch_id in batch_ids)

        assert wait_for(batches_finished, timeout=8)
        assert wait_for(lambda: aiapp.db.get_grading_job(
            shared_student_job.get_json()["jobId"])["status"] == "done", timeout=8)
        module_states = {
            item["key"]: item["status"] for item in get_json(
                client, "/api/admin/modules", admin).get_json()["modules"]
        }
        assert module_states["grading"] == "disabled"
        assert module_states["practice"] == "disabled"
        for module_key in ("grading", "practice", "review"):
            enabled = post_json(client, f"/api/admin/modules/{module_key}/state", {
                "state": "enabled",
            }, admin)
            assert enabled.status_code == 200, enabled.get_data(as_text=True)
        shared_done = wait_for_student_job(client, shared_student_job.get_json()["jobId"], student)
        assert shared_done["status"] == "done"
        data = aiapp.load_json(aiapp.ASSIGNMENTS_PATH, {"assignments": []})
        for batch_id in batch_ids:
            for essay in aiapp.find_assignment(data, batch_id)["essays"]:
                assert essay["attempt"] == 1
                assert essay["queuedAt"] and essay["startedAt"] and essay["finishedAt"]
                assert essay["durationSeconds"] is not None

        resp = post_json(client, "/api/assignments", {
            "title": "服务重启恢复测试", "prompt": "统一题目", "classId": "class",
        }, teacher)
        recovery_id = resp.get_json()["id"]
        assert post_json(client, f"/api/assignments/{recovery_id}/essays", {
            "essays": [{"studentName": "恢复学生", "text": "服务重启后运行中任务应恢复排队并继续完成。"}],
        }, teacher).status_code == 200
        recovery_reserved = aiapp.db.enqueue_assignment_jobs(recovery_id)
        recovery_job_id = recovery_reserved["jobs"][0]["jobId"]
        with aiapp.db.connection(write=True) as conn:
            conn.execute("UPDATE grading_jobs SET status='running',started_at='old' WHERE id=?",
                         (recovery_job_id,))
            conn.execute("UPDATE essays SET status='grading',started_at='old' WHERE job_id=?",
                         (recovery_job_id,))
        recovered_ids = aiapp.db.recover_grading_jobs()
        assert recovery_job_id in recovered_ids
        assert aiapp.db.get_grading_job(recovery_job_id)["status"] == "queued"
        ai_adapter.call_llm = fake_call_llm
        assert aiapp.schedule_grade_job(recovery_job_id)
        assert wait_for(lambda: aiapp.db.get_grading_job(recovery_job_id)["status"] == "done")

        resp = post_json(client, "/api/assignments", {
            "title": "正文版本测试", "prompt": "统一题目", "classId": "class",
        }, teacher)
        stale_id = resp.get_json()["id"]
        assert post_json(client, f"/api/assignments/{stale_id}/essays", {
            "essays": [{"studentName": "版本学生", "text": "这是等待批改时将被修改的原始作文正文。"}],
        }, teacher).status_code == 200
        stale_started = threading.Event()
        stale_release = threading.Event()

        def stale_call(messages, max_tokens=4096, temperature=0.6):
            stale_started.set()
            stale_release.wait(5)
            return fake_call_llm(messages, max_tokens, temperature)

        ai_adapter.call_llm = stale_call
        stale_start_response = post_json(
            client, f"/api/assignments/{stale_id}/grade", {}, teacher)
        assert stale_start_response.status_code == 200, stale_start_response.get_data(as_text=True)
        assert stale_started.wait(3)
        data = aiapp.load_json(aiapp.ASSIGNMENTS_PATH, {"assignments": []})
        stale_essay = aiapp.find_assignment(data, stale_id)["essays"][0]
        stale_essay["text"] = "正文已在旧任务返回前发生改变，因此旧结果必须被丢弃。"
        aiapp.save_json(aiapp.ASSIGNMENTS_PATH, data)
        stale_release.set()
        assert wait_for(lambda: aiapp.find_assignment(
            aiapp.load_json(aiapp.ASSIGNMENTS_PATH, {"assignments": []}), stale_id
        )["essays"][0]["status"] == "pending")

        transient_calls = 0

        def transient_once(messages, max_tokens=4096, temperature=0.6):
            nonlocal transient_calls
            transient_calls += 1
            if transient_calls == 1:
                raise aiapp.LLMRequestError("503", status_code=503, transient=True)
            return fake_call_llm(messages, max_tokens, temperature)

        ai_adapter.call_llm = transient_once
        assert aiapp.request_grade_result("题目", "要求", "", "", "足够长度的测试作文正文。")["score"] == 88
        assert transient_calls == 2

        format_calls = 0

        def malformed_once(messages, max_tokens=4096, temperature=0.6):
            nonlocal format_calls
            format_calls += 1
            return "不是 JSON" if format_calls == 1 else fake_call_llm(messages, max_tokens, temperature)

        ai_adapter.call_llm = malformed_once
        assert aiapp.request_grade_result("题目", "要求", "", "", "足够长度的测试作文正文。")["score"] == 88
        assert format_calls == 2

        permanent_calls = 0

        def permanent_error(messages, max_tokens=4096, temperature=0.6):
            nonlocal permanent_calls
            permanent_calls += 1
            raise aiapp.ApiKeyError("鉴权失败")

        ai_adapter.call_llm = permanent_error
        try:
            aiapp.request_grade_result("题目", "要求", "", "", "足够长度的测试作文正文。")
            raise AssertionError("永久错误不应被吞掉")
        except aiapp.ApiKeyError:
            pass
        assert permanent_calls == 1

    ai_adapter.call_llm = original_call_llm
    aiapp.GRADE_RETRY_DELAY = original_retry_delay
    os.environ.pop("ARTICLE_JUDGE_API_KEY", None)
    print("regression_test.py: OK")


if __name__ == "__main__":
    main()
