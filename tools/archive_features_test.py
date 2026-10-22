"""Isolated tests for revisions, trends, showcases, archives, and recovery."""

from __future__ import annotations

import datetime
import io
import json
import os
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

os.environ.setdefault("APP_ENV", "test")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import app as aiapp
from essay_platform.adapters import ai as ai_adapter
from essay_platform.services import recovery
from tools.regression_test import (
    get_json, login, post_json, put_json, reset_runtime, wait_for_student_job,
)


def insight_ai(messages, max_tokens=4096, temperature=0.6):
    assert "作文正文" not in messages[-1]["content"]
    payload = json.loads(messages[-1]["content"])
    assert payload["feedback"] and all("text" not in item for item in payload["feedback"])
    return json.dumps({
        "overall": "整体完成认真。",
        "commonStrengths": "选材贴近生活。",
        "commonProblems": "细节描写不足。",
        "teachingSuggestions": "安排一次细节扩写练习。",
    }, ensure_ascii=False)


def main():
    original_call = aiapp.call_llm
    with tempfile.TemporaryDirectory() as tmp:
        reset_runtime(tmp)
        client = aiapp.app.test_client()
        admin = login(client, "admin")
        teacher = login(client, "mr")
        student = login(client, "stu")
        ai_adapter.call_llm = aiapp.fake_call_llm if hasattr(aiapp, "fake_call_llm") else original_call

        assert put_json(client, "/api/classes/class", {
            "name": "一年级1班", "grade": "一年级", "klass": "1班",
        }, admin).status_code == 200
        created = post_json(client, "/api/assignments", {
            "title": "版本作文", "prompt": "写一件难忘的事", "classId": "class",
        }, teacher).get_json()
        aid = created["id"]
        assert post_json(client, f"/api/assignments/{aid}/publish", {}, teacher).status_code == 200

        def grading_ai(messages, max_tokens=4096, temperature=0.6):
            return json.dumps({
                "score": 88, "level": "良", "summary": "内容完整",
                "dimensions": {"立意与内容": 88, "结构与逻辑": 86, "语言表达": 87},
                "strengths": ["选材真实"], "weaknesses": ["细节偏少"],
                "suggestions": ["补充动作描写"], "modelSnippet": "示范句",
                "upgrade": "", "studentName": "",
            }, ensure_ascii=False)

        ai_adapter.call_llm = grading_ai
        first = post_json(client, "/api/student/grade", {
            "assignmentId": aid, "text": "这是第一稿作文，记录一件难忘的事情，正文长度足够用于批改。",
        }, student)
        assert first.status_code == 202, first.get_data(as_text=True)
        assert wait_for_student_job(client, first.get_json()["jobId"], student)["status"] == "pending_review"
        detail = get_json(client, f"/api/assignments/{aid}", teacher).get_json()
        eid = detail["essays"][0]["id"]
        assert post_json(client, f"/api/assignments/{aid}/essays/{eid}/publish-result", {}, teacher).status_code == 200
        first_history = get_json(client, "/api/student/history?type=assignment", student).get_json()["history"]
        assert len(first_history) == 1 and first_history[0]["score"] == 88

        assert post_json(client, f"/api/assignments/{aid}/essays/{eid}/return", {
            "comment": "请补充事情经过中的动作和心理描写。",
            "resubmitDeadlineAt": (datetime.datetime.now(datetime.timezone.utc)
                + datetime.timedelta(days=2)).isoformat().replace("+00:00", "Z"),
        }, teacher).status_code == 200
        returned = get_json(client, "/api/assignments", student).get_json()["assignments"][0]
        assert returned["submissionStatus"] == "resubmission_open"
        assert returned["canSubmit"] is True and "心理描写" in returned["returnComment"]

        second = post_json(client, "/api/student/grade", {
            "assignmentId": aid, "text": "这是第二稿作文，我补充了人物的动作、语言和心理描写，正文长度足够。",
        }, student)
        assert second.status_code == 202, second.get_data(as_text=True)
        assert wait_for_student_job(client, second.get_json()["jobId"], student)["status"] == "pending_review"
        timeline = get_json(client, f"/api/assignments/{aid}/essays/{eid}/versions", student)
        assert timeline.status_code == 200
        timeline_data = timeline.get_json()
        assert [item["versionNo"] for item in timeline_data["versions"]] == [1, 2]
        assert timeline_data["returns"][0]["resolvedAt"]
        assert "第一稿" in timeline_data["versions"][0]["text"]
        assert get_json(client, f"/api/assignments/{aid}/essays/{eid}/versions", admin).status_code == 404

        assert post_json(client, f"/api/assignments/{aid}/essays/{eid}/publish-result", {}, teacher).status_code == 200
        histories = get_json(client, "/api/student/history?type=assignment&pageSize=100", student).get_json()["history"]
        assert len(histories) == 2
        assert {item["publicationNo"] for item in histories} == {1, 2}
        trends = get_json(client, "/api/student/trends", student).get_json()
        assert len(trends["assignments"]) == 1 and trends["assignments"][0]["publicationNo"] == 2

        practice = post_json(client, "/api/student/grade", {
            "title": "自主练习", "text": "这是学生自主练习作文，老师端永远不应看到这条记录。",
        }, student)
        assert practice.status_code == 202
        assert wait_for_student_job(client, practice.get_json()["jobId"], student)["status"] == "done"
        trends = get_json(client, "/api/student/trends", student).get_json()
        assert len(trends["practice"]) == 1
        teacher_trend = get_json(client, "/api/teacher/students/stu/trends", teacher).get_json()
        assert len(teacher_trend["assignments"]) == 1 and "practice" not in teacher_trend

        saved = put_json(client, f"/api/assignments/{aid}/showcase", {
            "essayIds": [eid], "confirm": False,
        }, teacher)
        assert saved.status_code == 409
        saved = put_json(client, f"/api/assignments/{aid}/showcase", {
            "essayIds": [eid], "confirm": True,
        }, teacher)
        assert saved.status_code == 200
        showcase = get_json(client, f"/api/assignments/{aid}/showcase", student).get_json()["items"]
        assert len(showcase) == 1 and showcase[0]["studentName"] == "STU学生"
        assert "score" not in showcase[0] and "feedback" not in showcase[0]

        ai_adapter.call_llm = insight_ai
        insight = post_json(client, f"/api/assignments/{aid}/analytics/insight/generate", {}, teacher)
        assert insight.status_code == 200, insight.get_data(as_text=True)
        version = insight.get_json()["insight"]["version"]
        conflict = put_json(client, f"/api/assignments/{aid}/analytics/insight", {
            "version": version - 1,
            "insight": insight.get_json()["insight"]["insight"],
        }, teacher)
        assert conflict.status_code == 409

        completion = get_json(client, f"/api/assignments/{aid}/completion.xlsx", teacher)
        assert completion.status_code == 200
        sheet = load_workbook(io.BytesIO(completion.data), read_only=True).active
        completion_rows = list(sheet.iter_rows(values_only=True))
        assert completion_rows[0] == ("姓名", "学号", "提交状态", "等第")
        assert all(88 not in row for row in completion_rows)

        cloned = post_json(client, f"/api/assignments/{aid}/clone", {}, teacher)
        assert cloned.status_code == 200 and cloned.get_json()["lifecycleStatus"] == "draft"
        clone_id = cloned.get_json()["id"]
        assert client.delete(f"/api/assignments/{clone_id}", headers={"X-Token": teacher}).status_code == 200

        assert put_json(client, "/api/admin/school-settings", {
            "schoolName": "测试小学", "dailyTokenLimit": 1, "monthlyTokenLimit": 0,
        }, admin).status_code == 200
        aiapp.db.write_ai_usage(
            profile="article_judge", task_type="grading", model="fake",
            actor_username="stu", prompt_tokens=1,
        )
        blocked_insight = post_json(
            client, f"/api/assignments/{aid}/analytics/insight/generate",
            {"confirm": True}, teacher)
        assert blocked_insight.status_code == 503
        blocked_practice = post_json(client, "/api/student/grade", {
            "title": "限额测试", "text": "达到管理员限额后不应再创建新的自主练习批改任务。",
        }, student)
        assert blocked_practice.status_code == 503
        assert put_json(client, "/api/admin/school-settings", {
            "schoolName": "测试小学", "dailyTokenLimit": 0, "monthlyTokenLimit": 0,
        }, admin).status_code == 200

        modules = {item["key"]: item for item in get_json(
            client, "/api/admin/modules", admin).get_json()["modules"]}
        assert modules["ocr"]["implemented"] is True
        assert modules["ocr"]["status"] == "enabled"

        preview = post_json(client, "/api/admin/academic-years/rollover/preview", {}, admin)
        assert preview.status_code == 200
        preview_data = preview.get_json()
        assert preview_data["canConfirm"] is True
        rolled = post_json(client, "/api/admin/academic-years/rollover/confirm", {
            "targetStartYear": preview_data["target"]["startYear"],
            "classes": [item for item in preview_data["classes"] if not item["graduating"]],
        }, admin)
        assert rolled.status_code == 200, rolled.get_data(as_text=True)
        assert rolled.get_json()["createdClasses"][0]["teacherCount"] == 1
        current_class_id = rolled.get_json()["createdClasses"][0]["id"]
        teacher_classes = get_json(client, "/api/classes", teacher).get_json()["classes"]
        assert any(item["id"] == current_class_id for item in teacher_classes)
        archive = get_json(client, "/api/student/assignments/archive", student).get_json()
        assert not archive["current"] and archive["historical"][0]["id"] == aid
        archived_teacher_trend = get_json(client, "/api/teacher/students/stu/trends", teacher)
        assert archived_teacher_trend.status_code == 200
        assert len(archived_teacher_trend.get_json()["assignments"]) == 1
        assert post_json(client, "/api/student/grade", {
            "assignmentId": aid, "text": "历史作业不能再次提交，这段正文长度足够。",
        }, student).status_code in (403, 409)

        created_teacher = post_json(client, "/api/admin/users", {
            "role": "teacher", "username": "mr2", "employeeNo": "mr2",
            "display_name": "接收老师", "subject": "语文", "password": "Transfer123!",
            "classIds": [current_class_id],
        }, admin)
        assert created_teacher.status_code == 200, created_teacher.get_data(as_text=True)
        teacher2 = login(client, "mr2", "Transfer123!")
        transferred = post_json(client, "/api/admin/teachers/mr/transfer-assignments", {
            "targetUsername": "mr2", "assignmentIds": [aid],
        }, admin)
        assert transferred.status_code == 200
        assert get_json(client, f"/api/assignments/{aid}", teacher).status_code == 404
        assert get_json(client, f"/api/assignments/{aid}", teacher2).status_code == 200

        left_school = post_json(client, "/api/admin/students/stu/status", {
            "status": "transferred",
        }, admin)
        assert left_school.status_code == 200
        assert get_json(client, "/api/me", student).status_code == 401
        restored = post_json(client, "/api/admin/students/stu/status", {
            "status": "active", "classId": current_class_id,
        }, admin)
        assert restored.status_code == 200
        student = login(client, "stu", "stu-Secure123!")
        assert get_json(client, "/api/me", student).get_json()["schoolStatus"] == "active"

        backup_path = Path(aiapp.backup_data_snapshot("feature_test"))
        inspected = recovery.preview(aiapp.BACKUPS_DIR, backup_path.name)
        assert inspected["integrity"] == "ok" and inspected["counts"]["history"] >= 3

    ai_adapter.call_llm = original_call
    aiapp.db.reset_state_for_tests()
    print("archive_features_test.py: OK")


if __name__ == "__main__":
    main()
